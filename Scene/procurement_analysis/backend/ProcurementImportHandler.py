class ProcurementImportHandler:
    """POST /api/procurement/import - Parse Excel/CSV for procurement workbench."""

    def POST(self):
        _require_auth()
        _require_permission("scenes.use.procurement")
        web.header('Content-Type', 'application/json; charset=utf-8')
        try:
            import tempfile
            import uuid

            # web.py file upload
            upload_data = web.input(file={})
            file_item = upload_data.get('file')
            if not file_item or not file_item.filename:
                return json.dumps({"status": "error", "message": "No file uploaded"})

            filename = file_item.filename
            # Save to temp file
            temp_dir = tempfile.gettempdir()
            temp_path = os.path.join(temp_dir, f"procurement_{uuid.uuid4().hex[:8]}_{filename}")
            with open(temp_path, 'wb') as f:
                f.write(file_item.file.read())

            # Parse based on extension
            data = []
            try:
                if filename.lower().endswith('.csv'):
                    import csv
                    with open(temp_path, 'r', encoding='utf-8-sig') as f:
                        reader = csv.DictReader(f)
                        data = [row for row in reader]
                else:
                    # Excel: try data_only first, fallback to formula mode if bool conversion fails
                    from openpyxl import load_workbook
                    wb = None
                    try:
                        wb = load_workbook(temp_path, data_only=True)
                    except Exception as parse_err:
                        err_msg = str(parse_err)
                        if 'bool' in err_msg.lower() or 'converted' in err_msg.lower():
                            wb = load_workbook(temp_path, data_only=False)
                        else:
                            raise
                    ws = wb.active
                    rows = []
                    for row in ws.iter_rows(values_only=True):
                        cleaned = []
                        for cell in row:
                            if cell is None:
                                cleaned.append('')
                            elif isinstance(cell, bool):
                                cleaned.append('是' if cell else '否')
                            else:
                                cleaned.append(str(cell))
                        rows.append(cleaned)
                    if len(rows) >= 2:
                        headers = rows[0]
                        for r in rows[1:]:
                            if any(v.strip() for v in r):
                                data.append({h: v for h, v in zip(headers, r)})
            finally:
                try:
                    os.remove(temp_path)
                except Exception:
                    pass

            return json.dumps({
                "status": "success",
                "filename": filename,
                "data": data,
                "total_rows": len(data)
            }, ensure_ascii=False)
        except Exception as e:
            logger.error(f"[ProcurementImportHandler] Error: {e}")
            return json.dumps({"status": "error", "message": str(e)})
