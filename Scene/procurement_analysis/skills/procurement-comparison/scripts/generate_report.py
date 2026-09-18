#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate procurement comparison report in Excel or Word."""

import argparse
import json
import os
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    Document = None


def load_json_or_dict(value):
    if isinstance(value, dict):
        return value
    if Path(value).exists():
        with open(value, "r", encoding="utf-8") as f:
            return json.load(f)
    return json.loads(value)


def generate_excel(data: dict, output_path: str):
    if openpyxl is None:
        raise RuntimeError("openpyxl is required for Excel output")

    wb = openpyxl.Workbook()

    # Sheet 1: Summary
    ws = wb.active
    ws.title = "比价摘要"
    ws.merge_cells("A1:H1")
    ws["A1"] = "采购比价报告"
    ws["A1"].font = Font(size=18, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A3"] = f"报告日期：{datetime.now().strftime('%Y-%m-%d')}"
    ws["A4"] = f"参与供应商数：{data.get('quotes_count', 0)}"
    ws["A5"] = f"供应商：{'、'.join(data.get('suppliers', []))}"

    # Sheet 2: Comparison Matrix
    ws2 = wb.create_sheet("比价矩阵")
    headers = ["物料编码", "物料名称", "规格", "历史采购价", "最低价", "最高价", "建议目标价", "议价上限", "理想价", "建议"]
    ws2.append(headers)
    for cell in ws2[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")

    for row in data.get("comparison_matrix", []):
        ws2.append([
            row.get("物料编码", ""),
            row.get("物料名称", ""),
            row.get("规格", ""),
            row.get("历史采购价", 0),
            row.get("最低价", 0),
            row.get("最高价", 0),
            row.get("建议目标价", 0),
            row.get("议价上限", 0),
            row.get("理想价", 0),
            row.get("建议", "")
        ])

    # Add supplier columns
    supplier_col_start = 11
    for idx, supplier in enumerate(data.get("suppliers", [])):
        col = supplier_col_start + idx
        ws2.cell(row=1, column=col, value=f"{supplier}单价").font = Font(bold=True)
        ws2.cell(row=1, column=col + 1, value=f"{supplier}交期").font = Font(bold=True)
        for r_idx, row in enumerate(data.get("comparison_matrix", []), 2):
            quote_info = row.get("供应商报价", {}).get(supplier, {})
            ws2.cell(row=r_idx, column=col, value=quote_info.get("单价", ""))
            ws2.cell(row=r_idx, column=col + 1, value=quote_info.get("交期", ""))

    # Sheet 3: Supplier Scores
    ws3 = wb.create_sheet("供应商评分")
    score_headers = ["排名", "供应商", "价格得分", "交期得分", "质量得分", "服务得分", "综合得分", "平均报价", "平均交期"]
    ws3.append(score_headers)
    for cell in ws3[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")

    for score in data.get("supplier_scores", []):
        ws3.append([
            score.get("排名", 0),
            score.get("供应商", ""),
            score.get("价格得分", 0),
            score.get("交期得分", 0),
            score.get("质量得分", 0),
            score.get("服务得分", 0),
            score.get("综合得分", 0),
            score.get("平均报价", 0),
            score.get("平均交期", 0)
        ])

    # Sheet 4: Risks
    ws4 = wb.create_sheet("风险提示")
    ws4.append(["类型", "描述", "等级"])
    for cell in ws4[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")

    for risk in data.get("risks", []):
        ws4.append([risk.get("类型", ""), risk.get("描述", ""), risk.get("等级", "")])

    # Sheet 5: Recommendations
    ws5 = wb.create_sheet("采购建议")
    ws5.append(["物料", "建议供应商", "建议单价", "目标价", "理由"])
    for cell in ws5[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")

    for rec in data.get("recommendations", []):
        ws5.append([
            rec.get("物料", ""),
            rec.get("建议供应商", ""),
            rec.get("建议单价", 0),
            rec.get("目标价", 0),
            rec.get("理由", "")
        ])

    # Adjust column widths
    for sheet in wb.worksheets:
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 16

    wb.save(output_path)


def generate_word(data: dict, output_path: str):
    if Document is None:
        raise RuntimeError("python-docx is required for Word output")

    doc = Document()
    title = doc.add_heading("采购比价分析报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"报告日期：{datetime.now().strftime('%Y-%m-%d')}")
    doc.add_paragraph(f"参与供应商：{'、'.join(data.get('suppliers', []))}")

    doc.add_heading("一、比价矩阵", level=1)
    for row in data.get("comparison_matrix", []):
        doc.add_paragraph(
            f"{row.get('物料名称', '')}：最低价 {row.get('最低价', 0):.2f} 元，"
            f"历史价 {row.get('历史采购价', 0):.2f} 元，建议目标价 {row.get('建议目标价', 0):.2f} 元"
        )

    doc.add_heading("二、供应商评分", level=1)
    for score in data.get("supplier_scores", []):
        doc.add_paragraph(
            f"{score.get('供应商', '')}：综合得分 {score.get('综合得分', 0)} 分，"
            f"排名第 {score.get('排名', 0)}"
        )

    doc.add_heading("三、风险提示", level=1)
    for risk in data.get("risks", []):
        doc.add_paragraph(f"[{risk.get('等级', '')}]{risk.get('类型', '')}：{risk.get('描述', '')}")

    doc.add_heading("四、采购建议", level=1)
    for rec in data.get("recommendations", []):
        doc.add_paragraph(
            f"{rec.get('物料', '')}：建议选择 {rec.get('建议供应商', '')}，"
            f"目标价 {rec.get('目标价', 0):.2f} 元"
        )

    doc.save(output_path)


def generate_html(data: dict, output_path: str):
    """Generate a visual HTML procurement comparison dashboard."""
    matrix = data.get("comparison_matrix", [])
    scores = data.get("supplier_scores", [])
    risks = data.get("risks", [])

    suppliers = data.get("suppliers", [])
    materials = [row.get("物料名称", row.get("物料编码", f"物料{i+1}")) for i, row in enumerate(matrix)]

    # Price comparison data per material per supplier
    price_series = []
    for supplier in suppliers:
        values = []
        for row in matrix:
            quote = row.get("供应商报价", {}).get(supplier, {})
            values.append(quote.get("单价", 0))
        price_series.append({"name": supplier, "type": "bar", "data": values})

    # Score radar data
    radar_indicators = [
        {"name": "价格", "max": 100},
        {"name": "交期", "max": 100},
        {"name": "质量", "max": 100},
        {"name": "服务", "max": 100}
    ]
    radar_series = []
    colors = ["#3b82f6", "#10b981", "#f59e0b", "#ef4444"]
    for i, score in enumerate(scores):
        radar_series.append({
            "value": [
                score.get("价格得分", 0),
                score.get("交期得分", 0),
                score.get("质量得分", 0),
                score.get("服务得分", 0)
            ],
            "name": score.get("供应商", ""),
            "itemStyle": {"color": colors[i % len(colors)]}
        })

    matrix_rows = ""
    for row in matrix:
        supplier_cells = ""
        for supplier in suppliers:
            quote = row.get("供应商报价", {}).get(supplier, {})
            price = quote.get("单价", "-")
            delivery = quote.get("交期", "-")
            supplier_cells += f"<td class='p-3'>{price}<br><span class='text-xs text-gray-500'>{delivery}</span></td>"

        matrix_rows += f"""
        <tr class="border-b hover:bg-gray-50">
            <td class="p-3 font-medium">{row.get('物料名称', '')}</td>
            <td class="p-3">{row.get('历史采购价', 0):.2f}</td>
            <td class="p-3 text-green-600 font-semibold">{row.get('最低价', 0):.2f}</td>
            <td class="p-3 text-blue-600 font-semibold">{row.get('建议目标价', 0):.2f}</td>
            {supplier_cells}
            <td class="p-3 text-sm text-gray-600">{row.get('建议', '')}</td>
        </tr>
        """

    score_rows = ""
    for score in scores:
        score_rows += f"""
        <tr class="border-b">
            <td class="p-3 font-medium">{score.get('供应商', '')}</td>
            <td class="p-3">{score.get('价格得分', 0)}</td>
            <td class="p-3">{score.get('交期得分', 0)}</td>
            <td class="p-3">{score.get('质量得分', 0)}</td>
            <td class="p-3">{score.get('服务得分', 0)}</td>
            <td class="p-3 font-bold text-blue-600">{score.get('综合得分', 0)}</td>
            <td class="p-3">#{score.get('排名', 0)}</td>
        </tr>
        """

    risk_cards = ""
    risk_colors = {"高": "bg-red-50 text-red-700 border-red-200", "中": "bg-yellow-50 text-yellow-700 border-yellow-200", "低": "bg-green-50 text-green-700 border-green-200"}
    for risk in risks:
        level = risk.get("等级", "中")
        risk_cards += f"""
        <div class="p-4 rounded-lg border {risk_colors.get(level, 'bg-gray-50 text-gray-700 border-gray-200')}">
            <div class="font-semibold">[{level}] {risk.get('类型', '')}</div>
            <div class="text-sm mt-1">{risk.get('描述', '')}</div>
        </div>
        """

    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>采购比价报告</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://cdn.jsdelivr.net/npm/echarts@5.4.3/dist/echarts.min.js"></script>
</head>
<body class="bg-gray-50 p-6">
    <div class="max-w-7xl mx-auto space-y-6">
        <div class="bg-white rounded-xl shadow-sm p-6">
            <h1 class="text-2xl font-bold text-gray-800">采购比价报告</h1>
            <p class="text-gray-500 mt-1">参与供应商：{'、'.join(suppliers)}</p>
            <p class="text-gray-500">报告日期：{datetime.now().strftime('%Y-%m-%d')}</p>
        </div>

        <div class="grid grid-cols-1 md:grid-cols-2 gap-6">
            <div class="bg-white rounded-xl shadow-sm p-6">
                <h2 class="text-lg font-bold text-gray-800 mb-4">价格对比</h2>
                <div id="priceChart" style="height: 350px;"></div>
            </div>
            <div class="bg-white rounded-xl shadow-sm p-6">
                <h2 class="text-lg font-bold text-gray-800 mb-4">供应商评分雷达图</h2>
                <div id="radarChart" style="height: 350px;"></div>
            </div>
        </div>

        <div class="bg-white rounded-xl shadow-sm p-6">
            <h2 class="text-lg font-bold text-gray-800 mb-4">比价矩阵</h2>
            <div class="overflow-x-auto">
                <table class="w-full text-left border-collapse min-w-[700px]">
                    <thead>
                        <tr class="bg-gray-100">
                            <th class="p-3 rounded-tl-lg">物料</th>
                            <th class="p-3">历史采购价</th>
                            <th class="p-3">最低价</th>
                            <th class="p-3">建议目标价</th>
                            {''.join(f"<th class='p-3'>{s}</th>" for s in suppliers)}
                            <th class="p-3 rounded-tr-lg">建议</th>
                        </tr>
                    </thead>
                    <tbody>
                        {matrix_rows}
                    </tbody>
                </table>
            </div>
        </div>

        <div class="bg-white rounded-xl shadow-sm p-6">
            <h2 class="text-lg font-bold text-gray-800 mb-4">供应商评分卡</h2>
            <table class="w-full text-left border-collapse">
                <thead>
                    <tr class="bg-gray-100">
                        <th class="p-3 rounded-tl-lg">供应商</th>
                        <th class="p-3">价格得分</th>
                        <th class="p-3">交期得分</th>
                        <th class="p-3">质量得分</th>
                        <th class="p-3">服务得分</th>
                        <th class="p-3">综合得分</th>
                        <th class="p-3 rounded-tr-lg">排名</th>
                    </tr>
                </thead>
                <tbody>
                    {score_rows}
                </tbody>
            </table>
        </div>

        <div class="bg-white rounded-xl shadow-sm p-6">
            <h2 class="text-lg font-bold text-gray-800 mb-4">风险提示</h2>
            <div class="grid grid-cols-1 md:grid-cols-2 gap-4">
                {risk_cards if risk_cards else '<p class="text-gray-500">未发现明显风险</p>'}
            </div>
        </div>

        <div class="bg-white rounded-xl shadow-sm p-6">
            <h2 class="text-lg font-bold text-gray-800 mb-4">采购建议</h2>
            <div class="space-y-3">
                {''.join(f"<div class='p-4 rounded-lg border bg-gray-50'><span class='font-semibold text-gray-800'>{r.get('物料','')}：</span><span class='text-gray-700'>建议选择 {r.get('建议供应商','')}，目标价 {r.get('目标价',0):.2f} 元，理由：{r.get('理由','')}</span></div>" for r in data.get('recommendations', []))}
            </div>
            <div class="mt-5 p-4 bg-blue-50 rounded-lg">
                <h3 class="font-bold text-gray-800 mb-2">总体建议</h3>
                <ul class="text-sm text-gray-700 space-y-1">
                    <li>优先选择综合得分高的供应商</li>
                    <li>注意价格异常偏低的风险</li>
                    <li>争取有利的付款条件</li>
                    <li>建议和得分前两位的供应商进行进一步议价</li>
                </ul>
            </div>
        </div>
    </div>

    <script>
        const priceChart = echarts.init(document.getElementById('priceChart'));
        priceChart.setOption({{
            tooltip: {{ trigger: 'axis' }},
            legend: {{ data: {json.dumps(suppliers, ensure_ascii=False)} }},
            xAxis: {{ type: 'category', data: {json.dumps(materials, ensure_ascii=False)} }},
            yAxis: {{ type: 'value', name: '单价（元）' }},
            series: {json.dumps(price_series, ensure_ascii=False)}
        }});

        const radarChart = echarts.init(document.getElementById('radarChart'));
        radarChart.setOption({{
            tooltip: {{}},
            legend: {{ data: {json.dumps([s.get('供应商', '') for s in scores], ensure_ascii=False)} }},
            radar: {{
                indicator: {json.dumps(radar_indicators, ensure_ascii=False)},
                radius: '65%'
            }},
            series: [{{
                type: 'radar',
                data: {json.dumps(radar_series, ensure_ascii=False)}
            }}]
        }});

        window.addEventListener('resize', () => {{ priceChart.resize(); radarChart.resize(); }});
    </script>
</body>
</html>"""

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html)


def main():
    parser = argparse.ArgumentParser(description="Generate procurement comparison report")
    parser.add_argument("--comparison", required=True, help="Comparison result JSON file")
    parser.add_argument("--template", help="Template file path")
    parser.add_argument("--output", required=True, help="Output file path")
    parser.add_argument("--format", choices=["excel", "word", "html"], default="excel", help="Output format")
    args = parser.parse_args()

    data = load_json_or_dict(args.comparison)

    os.makedirs(os.path.dirname(args.output) or ".", exist_ok=True)

    if args.format == "excel":
        generate_excel(data, args.output)
    elif args.format == "word":
        generate_word(data, args.output)
    else:
        generate_html(data, args.output)

    print(f"Comparison report saved to: {args.output}")


if __name__ == "__main__":
    main()
