#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Create Excel template files for procurement-comparison skill."""

from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill


def create_history_price_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "历史采购价"
    headers = ["物料编码", "物料名称", "历史采购价", "最近采购日期", "历史供应商", "合格率(%)"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
    ws.append(["M001", "铝合金型材", 24.5, "2026-05-15", "供应商A", 98])
    ws.append(["M002", "螺丝M4", 0.48, "2026-05-20", "供应商B", 99])
    ws.append(["M003", "密封圈", 2.1, "2026-05-18", "供应商A", 97])
    for col in range(1, 7):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 16
    return wb


def create_supplier_profile_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "供应商档案"
    headers = ["供应商名称", "合作等级", "资质认证", "历史评分", "风险等级", "备注"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
    ws.append(["供应商A", "战略", "ISO9001,IATF16949", 92, "低", "长期合作"])
    ws.append(["供应商B", "合格", "ISO9001", 85, "低", "价格有优势"])
    ws.append(["供应商C", "临时", "ISO9001", 75, "中", "新供应商"])
    for col in range(1, 7):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18
    return wb


def create_report_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "比价报告模板"
    ws.merge_cells("A1:H1")
    ws["A1"] = "采购比价报告（模板）"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].alignment = openpyxl.styles.Alignment(horizontal="center")
    ws["A3"] = "本文件为模板，实际比价报告由 generate_report.py 自动生成"
    ws["A3"].font = Font(color="FF0000")
    return wb


def main():
    assets_dir = Path(__file__).parent.parent / "assets"
    assets_dir.mkdir(exist_ok=True)

    create_history_price_template().save(assets_dir / "history_price_template.xlsx")
    create_supplier_profile_template().save(assets_dir / "supplier_profile_template.xlsx")
    create_report_template().save(assets_dir / "comparison_report_template.xlsx")

    print(f"Templates created in: {assets_dir}")


if __name__ == "__main__":
    main()
