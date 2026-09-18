"""
排产功能API处理器
为生产计划-智能排产提供后端接口
"""

import json
import os
import tempfile
import uuid
import csv
from datetime import datetime, timedelta
from typing import List, Dict, Any, Optional

import web

from common.log import logger
from channel.web.web_channel_utils import _require_auth, _get_workspace_root


# ============================================================
# 数据持久化
# ============================================================

class SchedulingStore:
    """排产结果持久化存储"""

    def __init__(self):
        workspace = _get_workspace_root()
        self.store_dir = os.path.join(workspace, "scheduling_results")
        try:
            os.makedirs(self.store_dir, exist_ok=True)
        except OSError:
            # Fallback to temp directory if workspace is not writable
            import tempfile
            self.store_dir = os.path.join(tempfile.gettempdir(), "oneagent_scheduling_results")
            os.makedirs(self.store_dir, exist_ok=True)

    def _get_path(self, schedule_id: str) -> str:
        return os.path.join(self.store_dir, f"{schedule_id}.json")

    def save(self, data: Dict) -> str:
        """保存排产结果，返回schedule_id"""
        schedule_id = f"SCH_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:6]}"
        data["schedule_id"] = schedule_id
        data["created_at"] = datetime.now().isoformat()

        filepath = self._get_path(schedule_id)
        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2, default=str)
        return schedule_id

    def list(self, limit: int = 50) -> List[Dict]:
        """列出历史排产记录"""
        results = []
        for filename in sorted(os.listdir(self.store_dir), reverse=True):
            if filename.endswith(".json"):
                filepath = os.path.join(self.store_dir, filename)
                try:
                    with open(filepath, "r", encoding="utf-8") as f:
                        data = json.load(f)
                    # 构建显示名称
                    summary = data.get("summary", {})
                    order_count = summary.get("total_orders", "?")
                    on_time_rate = summary.get("on_time_rate", 0)
                    # 处理字符串格式的准时率（如 "100%"）
                    if isinstance(on_time_rate, str):
                        on_time_rate = float(on_time_rate.replace("%", "")) / 100
                    elif not isinstance(on_time_rate, float):
                        on_time_rate = 0
                    name = f"排产结果 - {order_count}个工单 - 准时率{int(on_time_rate*100)}%"
                    
                    # 计算总工期（从makespan或max_end_time）
                    makespan_days = summary.get("makespan_days", 0)
                    if not makespan_days and summary.get("max_end_time"):
                        try:
                            end_dt = datetime.fromisoformat(summary["max_end_time"].replace("Z", "+00:00"))
                            start_dt = datetime.fromisoformat(data.get("created_at", "").replace("Z", "+00:00"))
                            makespan_days = (end_dt - start_dt).days
                        except:
                            makespan_days = 0
                    
                    results.append({
                        "id": data.get("schedule_id"),
                        "name": name,
                        "created_at": data.get("created_at", "")[:19].replace("T", " "),
                        "summary": f"工单{order_count}个 | 准时{summary.get('on_time', 0)}个 | 延期{summary.get('tardy_orders', summary.get('delayed', 0))}个 | 总工期{makespan_days:.1f}天"
                    })
                    if len(results) >= limit:
                        break
                except Exception as e:
                    logger.error(f"[SchedulingStore] Error reading {filename}: {e}")
                    continue
        return results

    def get(self, schedule_id: str) -> Optional[Dict]:
        """获取指定排产记录"""
        filepath = self._get_path(schedule_id)
        if not os.path.exists(filepath):
            return None
        with open(filepath, "r", encoding="utf-8") as f:
            return json.load(f)


# ============================================================
# 数据解析工具
# ============================================================

def _parse_datetime(dt_str: str) -> datetime:
    """解析日期时间字符串"""
    formats = ["%Y-%m-%d", "%Y-%m-%d %H:%M", "%Y/%m/%d", "%Y/%m/%d %H:%M"]
    for fmt in formats:
        try:
            return datetime.strptime(dt_str, fmt)
        except ValueError:
            continue
    return datetime.now() + timedelta(days=7)


def _parse_process_route(route_input) -> List[Dict]:
    """解析工序路线，支持字符串或列表"""
    if not route_input:
        return [{"process_name": "默认工序", "time_per_unit": 60, "machine_id": None}]
    if isinstance(route_input, list):
        return route_input
    try:
        route = json.loads(route_input)
        if isinstance(route, list):
            return route
        return [route]
    except (json.JSONDecodeError, TypeError):
        # 尝试简单格式: "工序1:60|工序2:30"
        steps = []
        for part in str(route_input).split("|"):
            if ":" in part:
                name, time = part.split(":", 1)
                steps.append({"process_name": name.strip(), "time_per_unit": float(time), "machine_id": None})
        return steps if steps else [{"process_name": "默认工序", "time_per_unit": 60, "machine_id": None}]


def _parse_bom(bom_input) -> Dict[str, float]:
    """解析BOM，支持字符串或列表/字典"""
    if not bom_input:
        return {}
    if isinstance(bom_input, dict):
        return bom_input
    if isinstance(bom_input, list):
        # 列表格式: [{"material_name":"A","quantity":2.5}, ...]
        result = {}
        for item in bom_input:
            if isinstance(item, dict):
                name = item.get("material_name", item.get("name", ""))
                qty = item.get("quantity", item.get("qty", 0))
                if name:
                    result[name] = float(qty)
        return result
    try:
        parsed = json.loads(bom_input)
        if isinstance(parsed, dict):
            return parsed
        if isinstance(parsed, list):
            result = {}
            for item in parsed:
                if isinstance(item, dict):
                    name = item.get("material_name", item.get("name", ""))
                    qty = item.get("quantity", item.get("qty", 0))
                    if name:
                        result[name] = float(qty)
            return result
    except (json.JSONDecodeError, TypeError):
        pass
    # 尝试简单格式: "A:2.5,B:1.2"
    result = {}
    for part in str(bom_input).split(","):
        if ":" in part:
            name, qty = part.split(":", 1)
            result[name.strip()] = float(qty)
    return result


def _build_work_orders(data: List[Dict]) -> List[Any]:
    """从导入数据构建WorkOrder对象"""
    from agent.tools.scheduler.scheduler_engine import WorkOrder, ProcessStep

    orders = []
    for row in data:
        order_id = row.get("生产订单号", row.get("工单号", row.get("order_id", f"WO_{len(orders)+1}")))
        product = row.get("料号", row.get("品名", row.get("产品名称", row.get("product", "未知产品"))))
        product_code = row.get("料号", row.get("product_code", ""))
        quantity = int(row.get("计划数量", row.get("数量", row.get("quantity", 1))))
        due_date = _parse_datetime(row.get("需求时间", row.get("交期", row.get("due_date", ""))))
        priority_map = {"紧急": 1, "高": 2, "中": 3, "低": 4}
        priority_str = row.get("优先级", row.get("priority", "中"))
        priority = priority_map.get(priority_str, int(priority_str) if str(priority_str).isdigit() else 3)

        process_route = _parse_process_route(row.get("工序路线", row.get("process_route", "")))
        bom = _parse_bom(row.get("BOM物料清单", row.get("bom", "")))
        team = row.get("班组需求", row.get("team_required"))

        # 新模板字段
        raw_parent = row.get("父级", row.get("parent_id", None))
        if raw_parent == "0" or raw_parent == 0:
            parent_id = "0"
        elif raw_parent:
            parent_id = str(raw_parent).strip()
        else:
            parent_id = None
        sales_order_no = row.get("销售单号", row.get("sales_order_no", None))
        batch_no = row.get("批次", row.get("batch_no", None))
        batch_time = row.get("批次时间", row.get("batch_time", None))

        steps = [ProcessStep(
            process_name=s.get("process_name", s.get("工序名称", "默认工序")),
            time_per_unit=float(s.get("time_per_unit", s.get("time", s.get("工时（分钟/件）", 60))) or 60),
            machine_id=s.get("machine_id", s.get("设备编号", s.get("工作中心", None))),
            setup_time=float(s.get("setup_time", s.get("换线时间", 30.0)) or 30.0),
            sequence=int(s.get("sequence", s.get("行号", s.get("工序序号", 0))) or 0),
            mold_id=s.get("mold_id", s.get("模具", None))
        ) for s in process_route]
        order = WorkOrder(
            order_id=order_id,
            product=product,
            quantity=quantity,
            due_date=due_date,
            priority=priority,
            process_route=steps,
            bom=bom,
            team_required=team,
            parent_id=parent_id,
            sales_order_no=str(sales_order_no).strip() if sales_order_no else None,
            batch_no=str(batch_no).strip() if batch_no else None,
            batch_time=str(batch_time).strip() if batch_time else None,
            product_code=str(product_code).strip() if product_code else None
        )
        orders.append(order)
    return orders


def _build_machines(config: Dict) -> List[Any]:
    """构建设备资源"""
    from agent.tools.scheduler.scheduler_engine import Machine

    machines = []
    default_machines = [
        {"machine_id": "M001", "machine_name": "注塑机-1", "capacity_per_day": 500},
        {"machine_id": "M002", "machine_name": "注塑机-2", "capacity_per_day": 500},
        {"machine_id": "M003", "machine_name": "CNC-1", "capacity_per_day": 200},
        {"machine_id": "M004", "machine_name": "CNC-2", "capacity_per_day": 200},
        {"machine_id": "M005", "machine_name": "装配线-1", "capacity_per_day": 1000},
    ]

    machine_list = config.get("machines", default_machines)
    for m in machine_list:
        # 处理新模板的daily_capacity字段
        daily_capacity = m.get("daily_capacity")
        if daily_capacity and isinstance(daily_capacity, dict):
            m = dict(m)
            m["daily_capacity"] = daily_capacity
        machines.append(Machine(**{k: v for k, v in m.items() if k in ["machine_id","machine_name","capacity_per_day","available_start","available_end","setup_time_same","setup_time_diff","maintenance","setup_time_diff_mold","setup_time_same_mold_diff_product","daily_capacity"]}))
    return machines


def _build_teams(config: Dict) -> List[Any]:
    """构建班组资源"""
    from agent.tools.scheduler.scheduler_engine import Team

    teams = []
    default_teams = [
        {"team_name": "白班A组", "headcount": 12, "shift": "day"},
        {"team_name": "夜班B组", "headcount": 8, "shift": "night"},
    ]

    team_list = config.get("teams", default_teams)
    for t in team_list:
        teams.append(Team(**t))
    return teams


def _build_materials(config: Dict) -> List[Any]:
    """构建物料数据"""
    from agent.tools.scheduler.scheduler_engine import Material

    materials = []
    default_materials = [
        {"material_name": "原材料A", "stock": 1000, "on_the_way": 500, "lead_time": 3},
        {"material_name": "原材料B", "stock": 500, "on_the_way": 0, "lead_time": 5},
    ]

    mat_list = config.get("materials", default_materials)
    for m in mat_list:
        materials.append(Material(**m))
    return materials


# ============================================================
# API Handlers
# ============================================================

class SchedulingScheduleHandler:
    """POST /api/scheduling/schedule - 执行排产"""

    def POST(self):
        _require_auth()
        web.header('Content-Type', 'application/json; charset=utf-8')
        try:
            body = json.loads(web.data() or b"{}")
            orders_data = body.get("orders", [])
            config = body.get("config", {})
            params = body.get("params", {})

            if not orders_data:
                return json.dumps({"status": "error", "message": "No orders provided"}, ensure_ascii=False)

            # 构建数据模型
            orders = _build_work_orders(orders_data)
            machines = _build_machines(config)
            teams = _build_teams(config)
            materials = _build_materials(config)

            # 执行排产
            from agent.tools.scheduler.scheduler_engine import AdvancedScheduler
            scheduler = AdvancedScheduler()
            scheduler.set_resources(machines, teams)
            scheduler.set_materials(materials)

            # 设置参数
            scheduler.schedule_mode = params.get("mode", "forward")
            scheduler.objective = params.get("objective", "tardiness")
            scheduler.iterations = params.get("iterations", 500)

            result = scheduler.schedule(orders)

            # 生成甘特图HTML
            from agent.tools.scheduler.gantt_generator import GanttGenerator
            gantt_gen = GanttGenerator()
            gantt_html = gantt_gen.generate(
                result["schedule"],
                title="正向排产 - 生产排产计划" if params.get("mode") == "forward" else "逆向排产 - 生产排产计划",
                order_gantt=result.get("order_gantt", [])
            )

            # 保存结果
            store = SchedulingStore()
            schedule_id = store.save({
                "params": params,
                "config": config,
                "orders": orders_data,
                "result": result,
                "gantt_html": gantt_html
            })

            return json.dumps({
                "status": "success",
                "schedule_id": schedule_id,
                "summary": result["summary"],
                "alerts": result["alerts"],
                "bottleneck": result["bottleneck"],
                "material_plan": result["material_plan"],
                "gantt_html": gantt_html
            }, ensure_ascii=False, default=str)

        except Exception as e:
            logger.error(f"[SchedulingScheduleHandler] Error: {e}")
            return json.dumps({"status": "error", "message": str(e)}, ensure_ascii=False)


class SchedulingGanttHandler:
    """POST /api/scheduling/gantt - 生成甘特图"""

    def POST(self):
        _require_auth()
        web.header('Content-Type', 'application/json; charset=utf-8')
        try:
            body = json.loads(web.data() or b"{}")
            schedule = body.get("schedule", [])
            title = body.get("title", "生产排产甘特图")

            from agent.tools.scheduler.gantt_generator import GanttGenerator
            gantt_gen = GanttGenerator()
            html = gantt_gen.generate(schedule, title)

            return json.dumps({
                "status": "success",
                "html": html
            }, ensure_ascii=False)

        except Exception as e:
            logger.error(f"[SchedulingGanttHandler] Error: {e}")
            return json.dumps({"status": "error", "message": str(e)}, ensure_ascii=False)


class SchedulingMaterialCheckHandler:
    """POST /api/scheduling/material-check - 物料齐套检查"""

    def POST(self):
        _require_auth()
        web.header('Content-Type', 'application/json; charset=utf-8')
        try:
            body = json.loads(web.data() or b"{}")
            orders_data = body.get("orders", [])
            materials_data = body.get("materials", [])

            if not orders_data:
                return json.dumps({"status": "error", "message": "No orders provided"}, ensure_ascii=False)

            orders = _build_work_orders(orders_data)
            materials = _build_materials({"materials": materials_data})

            from agent.tools.scheduler.scheduler_engine import MaterialChecker
            checker = MaterialChecker(materials)

            results = []
            for order in orders:
                ready, result = checker.check_availability(order)
                results.append(result)

            return json.dumps({
                "status": "success",
                "results": results
            }, ensure_ascii=False, default=str)

        except Exception as e:
            logger.error(f"[SchedulingMaterialCheckHandler] Error: {e}")
            return json.dumps({"status": "error", "message": str(e)}, ensure_ascii=False)


class SchedulingBottleneckHandler:
    """POST /api/scheduling/bottleneck - 瓶颈分析"""

    def POST(self):
        _require_auth()
        web.header('Content-Type', 'application/json; charset=utf-8')
        try:
            body = json.loads(web.data() or b"{}")
            schedule = body.get("schedule", [])

            if not schedule:
                return json.dumps({"status": "error", "message": "No schedule data provided"}, ensure_ascii=False)

            # 计算设备负载
            machine_loads = {}
            machine_orders = {}
            for item in schedule:
                mid = item.get("machine_id", "未知")
                start = datetime.fromisoformat(item["start_time"])
                end = datetime.fromisoformat(item["end_time"])
                duration = (end - start).total_seconds() / 3600
                machine_loads[mid] = machine_loads.get(mid, 0) + duration
                if item.get("order_id") not in machine_orders.get(mid, []):
                    machine_orders.setdefault(mid, []).append(item.get("order_id"))

            if not machine_loads:
                return json.dumps({"status": "success", "bottleneck": {}}, ensure_ascii=False)

            bottleneck = max(machine_loads, key=machine_loads.get)
            total_available = 12 * 7  # 一周
            utilization = min(100, (machine_loads[bottleneck] / total_available) * 100)

            return json.dumps({
                "status": "success",
                "bottleneck": {
                    "bottleneck_machine": bottleneck,
                    "bottleneck_load_hours": round(machine_loads[bottleneck], 2),
                    "utilization_percent": round(utilization, 1),
                    "affected_orders": machine_orders.get(bottleneck, []),
                    "all_machines": {k: round(v, 2) for k, v in machine_loads.items()}
                }
            }, ensure_ascii=False)

        except Exception as e:
            logger.error(f"[SchedulingBottleneckHandler] Error: {e}")
            return json.dumps({"status": "error", "message": str(e)}, ensure_ascii=False)


class SchedulingWhatIfHandler:
    """POST /api/scheduling/what-if - 急单插单影响分析"""

    def POST(self):
        _require_auth()
        web.header('Content-Type', 'application/json; charset=utf-8')
        try:
            body = json.loads(web.data() or b"{}")
            orders_data = body.get("orders", [])
            rush_order_data = body.get("rush_order")
            config = body.get("config", {})

            if not orders_data or not rush_order_data:
                return json.dumps({"status": "error", "message": "Orders and rush_order required"}, ensure_ascii=False)

            orders = _build_work_orders(orders_data)
            rush_orders = _build_work_orders([rush_order_data])
            rush_order = rush_orders[0] if rush_orders else None

            machines = _build_machines(config)
            teams = _build_teams(config)
            materials = _build_materials(config)

            from agent.tools.scheduler.scheduler_engine import AdvancedScheduler
            scheduler = AdvancedScheduler()
            scheduler.set_resources(machines, teams)
            scheduler.set_materials(materials)

            result = scheduler.what_if_analysis(orders, rush_order)

            return json.dumps({
                "status": "success",
                **result
            }, ensure_ascii=False, default=str)

        except Exception as e:
            logger.error(f"[SchedulingWhatIfHandler] Error: {e}")
            return json.dumps({"status": "error", "message": str(e)}, ensure_ascii=False)


class SchedulingImportHandler:
    """POST /api/scheduling/import - 导入CSV/Excel排产数据"""

    def POST(self):
        _require_auth()
        web.header('Content-Type', 'application/json; charset=utf-8')
        try:
            upload_data = web.input(file={})
            file_item = upload_data.get('file')
            if not file_item or not file_item.filename:
                return json.dumps({"status": "error", "message": "No file uploaded"}, ensure_ascii=False)

            filename = file_item.filename
            temp_dir = tempfile.gettempdir()
            temp_path = os.path.join(temp_dir, f"scheduling_{uuid.uuid4().hex[:8]}_{filename}")
            with open(temp_path, 'wb') as f:
                f.write(file_item.file.read())

            data = []
            try:
                if filename.lower().endswith('.csv'):
                    with open(temp_path, 'r', encoding='utf-8-sig') as f:
                        reader = csv.DictReader(f)
                        data = [row for row in reader]
                else:
                    from openpyxl import load_workbook
                    wb = load_workbook(temp_path, data_only=True)
                    sheets_data = {}
                    for sheet_name in wb.sheetnames:
                        ws = wb[sheet_name]
                        rows = []
                        for row in ws.iter_rows(values_only=True):
                            rows.append([str(cell) if cell is not None else '' for cell in row])
                        if len(rows) >= 2:
                            headers = rows[0]
                            sheet_data = []
                            for r in rows[1:]:
                                if any(v.strip() for v in r):
                                    sheet_data.append({h: v for h, v in zip(headers, r)})
                            sheets_data[sheet_name] = sheet_data
                    data = sheets_data
            finally:
                try:
                    os.remove(temp_path)
                except Exception:
                    pass

            return json.dumps({
                "status": "success",
                "filename": filename,
                "data": data,
                "total_rows": len(data) if isinstance(data, list) else sum(len(v) for v in data.values())
            }, ensure_ascii=False)

        except Exception as e:
            logger.error(f"[SchedulingImportHandler] Error: {e}")
            return json.dumps({"status": "error", "message": str(e)}, ensure_ascii=False)


class SchedulingHistoryHandler:
    """GET /api/scheduling/history - 获取历史排产记录"""

    def GET(self):
        _require_auth()
        web.header('Content-Type', 'application/json; charset=utf-8')
        try:
            params = web.input(limit='50')
            limit = int(params.limit)

            store = SchedulingStore()
            history = store.list(limit=limit)

            return json.dumps({
                "status": "success",
                "records": history,
                "total": len(history)
            }, ensure_ascii=False, default=str)

        except Exception as e:
            logger.error(f"[SchedulingHistoryHandler] Error: {e}")
            return json.dumps({"status": "error", "message": str(e)}, ensure_ascii=False)


class SchedulingHistoryDetailHandler:
    """GET /api/scheduling/history/(.*) - 获取指定排产记录详情"""

    def GET(self, schedule_id: str):
        _require_auth()
        web.header('Content-Type', 'application/json; charset=utf-8')
        try:
            store = SchedulingStore()
            data = store.get(schedule_id)

            if not data:
                return json.dumps({"status": "error", "message": "Schedule not found"}, ensure_ascii=False)

            return json.dumps({
                "status": "success",
                "record": data
            }, ensure_ascii=False, default=str)

        except Exception as e:
            logger.error(f"[SchedulingHistoryDetailHandler] Error: {e}")
            return json.dumps({"status": "error", "message": str(e)}, ensure_ascii=False)


class SchedulingBOMTreeHandler:
    """POST /api/scheduling/bom-tree - BOM树排程（融合方案）"""

    def POST(self):
        _require_auth()
        web.header('Content-Type', 'application/json; charset=utf-8')
        try:
            body = json.loads(web.data() or b"{}")
            orders_data = body.get("orders", [])
            config = body.get("config", {})
            params = body.get("params", {})
            bom_data = body.get("bom", {})  # BOM数据 {产品名: [{material_name, quantity, item_type}, ...]}
            process_data = body.get("process_db", {})  # 工序数据库

            if not orders_data:
                return json.dumps({"status": "error", "message": "No orders provided"}, ensure_ascii=False)

            # 构建数据模型
            orders = _build_work_orders(orders_data)
            machines = _build_machines(config)
            teams = _build_teams(config)
            materials = _build_materials(config)

            # 转换BOM数据
            from agent.tools.scheduler.bom_tree import BOMItem
            bom_db = {}
            for product, items in bom_data.items():
                bom_db[product] = []
                for item in items:
                    bom_db[product].append(BOMItem(
                        material_name=item.get("material_name", item.get("品名", item.get("物料名称", ""))),
                        material_code=item.get("material_code", item.get("料号", "")),
                        quantity=float(item.get("quantity", item.get("用量", 0))),
                        item_type=item.get("item_type", item.get("类型", "raw_material")),
                        unit=item.get("unit", item.get("单位", "个")),
                        base_qty=float(item.get("base_qty", item.get("底数", 1))) or 1.0
                    ))

            # 转换工序数据库
            process_db = {}
            for product, steps in process_data.items():
                from agent.tools.scheduler.scheduler_engine import ProcessStep
                process_db[product] = [
                    ProcessStep(
                        process_name=s.get("process_name", s.get("工序名称", "默认工序")),
                        time_per_unit=float(s.get("time_per_unit", s.get("耗时", 60))),
                        machine_id=s.get("machine_id", s.get("设备编号", None)),
                        mold_id=s.get("mold_id", s.get("模具", None)),
                        setup_time=float(s.get("setup_time", s.get("换线时间", 30))),
                        sequence=int(s.get("sequence", s.get("行号", s.get("工序序号", 0)))) or 0
                    )
                    for s in steps
                ]

            # 执行BOM树排程
            import time
            from agent.tools.scheduler.scheduler_engine import AdvancedScheduler
            scheduler = AdvancedScheduler()
            scheduler.set_resources(machines, teams)
            scheduler.set_materials(materials)

            # 设置参数
            scheduler.schedule_mode = params.get("mode", "backward")  # BOM树排程默认倒排
            scheduler.objective = params.get("objective", "tardiness")
            scheduler.iterations = params.get("iterations", 500)

            transfer_time = float(params.get("transfer_time_hours", 5.0))

            logger.info(f"[BOMTree] Starting schedule_with_bom_tree: {len(orders)} orders, {len(process_db)} process routes, transfer={transfer_time}")
            t0 = time.time()
            result = scheduler.schedule_with_bom_tree(
                orders, bom_db, process_db, transfer_time
            )
            t1 = time.time()
            logger.info(f"[BOMTree] schedule_with_bom_tree completed in {t1-t0:.2f}s, summary={result.get('summary')}")

            # 生成统一的甘特图HTML（包含设备视图、订单甘特图、排产明细表）
            from agent.tools.scheduler.gantt_generator import GanttGenerator
            gantt_gen = GanttGenerator()
            gantt_html = gantt_gen.generate(
                result["schedule"],
                "BOM树排程 - 生产排产计划",
                order_gantt=result.get("order_gantt", [])
            )

            # 保存结果
            store = SchedulingStore()
            schedule_id = store.save({
                "params": params,
                "config": config,
                "orders": orders_data,
                "result": result,
                "gantt_html": gantt_html
            })

            return json.dumps({
                "status": "success",
                "schedule_id": schedule_id,
                "summary": result["summary"],
                "alerts": result["alerts"],
                "material_kit_result": result["material_kit_result"],
                "gantt_html": gantt_html
            }, ensure_ascii=False, default=str)

        except Exception as e:
            logger.error(f"[SchedulingBOMTreeHandler] Error: {e}")
            import traceback
            logger.error(traceback.format_exc())
            return json.dumps({"status": "error", "message": str(e)}, ensure_ascii=False)
