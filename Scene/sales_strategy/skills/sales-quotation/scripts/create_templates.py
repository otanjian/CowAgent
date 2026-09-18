#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Create Excel template files for sales-quotation skill."""

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


def create_bom_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM清单"
    headers = ["物料编码", "物料名称", "用量", "单位", "损耗率(%)", "单价", "备注"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
    ws.append(["M001", "铝合金型材", 1.2, "kg", 5, 25.0, "主料"])
    ws.append(["M002", "螺丝M4", 4, "个", 2, 0.5, "标准件"])
    ws.append(["M003", "密封圈", 1, "个", 3, 2.0, ""] )
    for col in range(1, 8):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15
    return wb


def create_process_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "工艺路线"
    headers = ["工序", "设备", "工时(小时)", "人工费率(元/小时)", "制造费用率(元/小时)", "备注"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
    ws.append(["下料", "锯床", 0.5, 35.0, 20.0, ""])
    ws.append(["CNC加工", "CNC-1", 1.5, 45.0, 60.0, ""])
    ws.append(["表面处理", "氧化线", 0.8, 30.0, 40.0, ""])
    ws.append(["装配", "装配线", 0.6, 30.0, 25.0, ""])
    for col in range(1, 7):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18
    return wb


def create_inventory_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "库存数据"
    headers = ["物料编码", "物料名称", "需求数量", "可用库存", "采购周期(天)", "备注"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
    ws.append(["M001", "铝合金型材", 1.2, 800, 7, "单件需求，将乘以订单数量"])
    ws.append(["M002", "螺丝M4", 4, 5000, 3, "单件需求，将乘以订单数量"])
    ws.append(["M003", "密封圈", 1, 1200, 5, "单件需求，将乘以订单数量"])
    for col in range(1, 7):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 16
    return wb


def create_quotation_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "报价单模板"
    ws.merge_cells("A1:F1")
    ws["A1"] = "销售报价单（模板）"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A3"] = "本文件为模板，实际报价单由 generate_quotation.py 自动生成"
    ws["A3"].font = Font(color="FF0000")
    return wb


def create_rates_template():
    return {
        "管理费率": 0.1,
        "币种": "CNY",
        "汇率": 1.0
    }


def create_capacity_template():
    return {
        "日产能": 500,
        "当前负荷": 200,
        "缓冲天数": 2
    }


def main():
    assets_dir = Path(__file__).parent.parent / "assets"
    assets_dir.mkdir(exist_ok=True)

    create_bom_template().save(assets_dir / "bom_template.xlsx")
    create_process_template().save(assets_dir / "process_template.xlsx")
    create_inventory_template().save(assets_dir / "inventory_template.xlsx")
    create_quotation_template().save(assets_dir / "quotation_template.xlsx")

    with open(assets_dir / "rates_template.json", "w", encoding="utf-8") as f:
        json.dump(create_rates_template(), f, ensure_ascii=False, indent=2)

    with open(assets_dir / "capacity_template.json", "w", encoding="utf-8") as f:
        json.dump(create_capacity_template(), f, ensure_ascii=False, indent=2)

    print(f"Templates created in: {assets_dir}")


if __name__ == "__main__":
    main()
