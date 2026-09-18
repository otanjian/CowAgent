"""气袋排产 API Handler。

路由（统一前缀，由 action 分发）：
  POST /api/airbag-scheduling/parse              解析预览（场景权限）
  POST /api/airbag-scheduling/template-content  模板各 sheet 内容预览（场景权限）
  GET  /api/airbag-scheduling/demands           预估出货量表信息与内容预览（场景权限）
  POST /api/airbag-scheduling/demands           上传预估出货量表（场景权限，排产1 数据来源）
  GET  /api/airbag-scheduling/config             规则参数读取（场景权限）
  PUT  /api/airbag-scheduling/config             规则参数保存（+scheduling.config.manage）
  GET  /api/airbag-scheduling/capacity           产能配置表读取（场景权限）
  PUT  /api/airbag-scheduling/capacity           产能配置表保存（+scheduling.config.manage）
  GET  /api/airbag-scheduling/teams              班组产线配置表读取（场景权限）
  PUT  /api/airbag-scheduling/teams              班组产线配置表保存（+scheduling.config.manage）
  GET  /api/airbag-scheduling/merge-rules        拼线规则表读取（场景权限）
  PUT  /api/airbag-scheduling/merge-rules        拼线规则表保存（+scheduling.config.manage）
  POST /api/airbag-scheduling/sync-from-excel    从上传 Excel 同步产能/班组/拼线（+scheduling.config.manage）
  POST /api/airbag-scheduling/run                执行排产（场景权限）
  GET  /api/airbag-scheduling/history            历史记录（场景权限）
  GET  /api/airbag-scheduling/template           模板信息（场景权限）
  POST /api/airbag-scheduling/template           更换上传模板（+scheduling.config.manage）
  POST /api/airbag-scheduling/template/reset     恢复内置模板（+scheduling.config.manage）
"""

import base64
import json
import os
import re
import subprocess
import sys
import time
from datetime import datetime, timedelta

import web

from channel.web.web_channel_utils import (
    _get_workspace_root,
    _require_auth,
    _require_permission,
)
from common.log import logger

# 场景入口权限
SCENE_PERMISSION = "scenes.use.production"
# 配置写操作权限
CONFIG_MANAGE_PERMISSION = "scheduling.config.manage"

# 技能名（场景配置中的 skill_name）
SKILL_NAME = "pmc-scheduler-hmt-qd"
# 内置模板文件名
BUILTIN_TEMPLATE_NAME = "生产排产模板.xlsx"

# sheet 名候选（兼容新模板：去掉无用 sheet 后仅保留单个 sheet1）
S2_SHEET_CANDIDATES = ("排产2", "SHEET1", "Sheet1", "sheet1")
DEMAND_SHEET_CANDIDATES = ("排产1", "SHEET1", "Sheet1", "sheet1", "出货量")
# 排产2 表（缝纫周计划）仅读取/展示至 BG 列（59），其后多余内容忽略，避免超大列导致卡顿
S2_PREVIEW_MAX_COLS = 59

# 配置存储目录（相对 workspace）：{workspace}/.one/scheduling/airbag/
CONFIG_SUBDIR = os.path.join(".one", "scheduling", "airbag")
# 排产结果输出目录（相对 workspace）：{workspace}/scheduling_results/airbag/
RESULT_SUBDIR = os.path.join("scheduling_results", "airbag")

# 配置文件名与引擎 CLI 参数映射
CONFIG_FILES = {
    "rules": ("rules.json", "--config"),
    "capacity": ("capacity.json", "--capacity"),
    "teams": ("teams.json", "--teams"),
    "merge_rules": ("merge_rules.json", "--merge"),
}

# 排产规则默认参数（与技能引擎 DEFAULT_RULES 保持一致，文件不存在时返回给前端）
DEFAULT_RULES = {
    "version": 1,
    "period_count": 2,
    "max_priority": 2,
    "shebian_keyword": "设变",
    "stockup_keywords": ["2周", "3周", "4周"],
    "stockup_trigger_keywords": ["提前备货"],
    "stockup_start_on_monday": True,
    "night_divisor": 14,
    "night_basic_enabled": True,
    "night_enabled": True,
    "early_col_parity": "odd",
}


# ---------------------------------------------------------------------------
# 内部工具
# ---------------------------------------------------------------------------
def _project_root():
    """项目根目录（handlers → web → channel → root 四层 dirname）。"""
    return os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))


def _scheduling_config_dir():
    """配置存储目录 {workspace}/.one/scheduling/airbag/，不存在时创建。"""
    d = os.path.join(_get_workspace_root(), CONFIG_SUBDIR)
    os.makedirs(d, exist_ok=True)
    return d


def _result_dir():
    """排产结果输出目录 {workspace}/scheduling_results/airbag/，不存在时创建。"""
    d = os.path.join(_get_workspace_root(), RESULT_SUBDIR)
    os.makedirs(d, exist_ok=True)
    return d


def _resolve_skill_dir():
    """定位技能目录：租户/工作区 skills 优先，其次项目根 skills。"""
    candidates = [
        os.path.join(_get_workspace_root(), "skills", SKILL_NAME),
        os.path.join(_project_root(), "skills", SKILL_NAME),
    ]
    for c in candidates:
        if os.path.isdir(c):
            return c
    return None


def _safe_filename(name):
    """文件名消毒：仅保留字母数字中文字符、点、横线、下划线。"""
    name = os.path.basename(name or "")
    name = re.sub(r"[^\w.\-\u4e00-\u9fa5]", "_", name)
    return name or "template.xlsx"


def _save_uploaded_file(body, key="excel"):
    """从请求体取出 base64 上传文件并保存到 {workspace}/tmp/airbag_scheduling/。

    返回 (文件绝对路径, 错误消息)；文件不存在/未上传时返回 (None, None)。
    """
    files = body.get("files") or {}
    info = files.get(key) or {}
    if not info:
        return None, None
    filename = _safe_filename(info.get("filename", "template.xlsx"))
    content = info.get("content", "")
    is_base64 = info.get("is_base64", False)
    if not content:
        return None, "上传文件内容为空"

    upload_dir = os.path.join(_get_workspace_root(), "tmp", "airbag_scheduling")
    os.makedirs(upload_dir, exist_ok=True)
    path = os.path.join(upload_dir, f"{int(time.time() * 1000)}_{filename}")
    try:
        if is_base64:
            with open(path, "wb") as f:
                f.write(base64.b64decode(content))
        else:
            with open(path, "w", encoding="utf-8") as f:
                f.write(content)
    except Exception as e:  # noqa: BLE001
        logger.error(f"[AirbagScheduling] 保存上传文件失败: {e}")
        return None, f"上传文件保存失败: {e}"
    return path, None


def _resolve_template(use_uploaded=True):
    """返回 (模板绝对路径, 来源) 。

    来源：uploaded（用户上次上传）/ builtin（技能内置）/ None（未找到）。
    """
    if use_uploaded:
        uploaded = os.path.join(_scheduling_config_dir(), "template.xlsx")
        if os.path.isfile(uploaded):
            return uploaded, "uploaded"
    skill_dir = _resolve_skill_dir()
    if skill_dir:
        builtin = os.path.join(skill_dir, "templates", BUILTIN_TEMPLATE_NAME)
        if os.path.isfile(builtin):
            return builtin, "builtin"
    return None, None


def _template_info():
    """当前模板信息：来源、上传时间、文件名。"""
    uploaded = os.path.join(_scheduling_config_dir(), "template.xlsx")
    if os.path.isfile(uploaded):
        return {
            "source": "uploaded",
            "uploaded_at": datetime.fromtimestamp(os.path.getmtime(uploaded)).strftime("%Y-%m-%d %H:%M:%S"),
            "filename": "template.xlsx",
        }
    skill_dir = _resolve_skill_dir()
    builtin = os.path.join(skill_dir, "templates", BUILTIN_TEMPLATE_NAME) if skill_dir else ""
    if os.path.isfile(builtin):
        return {
            "source": "builtin",
            "uploaded_at": None,
            "filename": BUILTIN_TEMPLATE_NAME,
        }
    return {"source": "missing", "uploaded_at": None, "filename": ""}


def _demands_path():
    """预估出货量表路径 {workspace}/.one/scheduling/airbag/demands.xlsx；不存在返回 None。"""
    p = os.path.join(_scheduling_config_dir(), "demands.xlsx")
    return p if os.path.isfile(p) else None


def _demands_info():
    """预估出货量表信息：是否存在、上传时间、文件名。"""
    p = _demands_path()
    if not p:
        return {"exists": False, "uploaded_at": None, "filename": ""}
    return {
        "exists": True,
        "uploaded_at": datetime.fromtimestamp(os.path.getmtime(p)).strftime("%Y-%m-%d %H:%M:%S"),
        "filename": "demands.xlsx",
    }


def _normalize_sap_zeros(filepath):
    """去掉 Excel 内 SAP 物料号列（B 列）的前导零，使界面展示更简洁（如 000000005000001256 → 5000001256）。

    仅处理纯数字且以 0 开头的单元格；失败时静默跳过，不影响导入。
    """
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, keep_links=False)
        changed = 0
        for ws in wb.worksheets:
            for row in range(2, ws.max_row + 1):
                v = ws.cell(row=row, column=2).value
                if v is None:
                    continue
                s = str(v).strip()
                if s.isdigit() and s.startswith("0"):
                    ws.cell(row=row, column=2).value = s.lstrip("0")
                    changed += 1
        if changed:
            wb.save(filepath)
    except Exception:  # noqa: BLE001 规范化非关键步骤，失败不影响导入
        pass


def _to_date(v):
    """将表头单元格转 date；支持 datetime、8位整数/字符串、YYYY/M/D、带时分秒日期串、Excel 日期序列号。

    客户文件周列表头日期格式不固定（可能是 '20260814'、'2026/8/14'、'2026-08-14 00:00:00'
    或 Excel 日期序列号），必须一并识别。
    """
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, (int, float)):
        n = int(v) if abs(float(v) - round(float(v))) < 1e-6 else None
        if n is None:
            return None
        s = str(n).strip()
        if len(s) == 8 and s.isdigit():
            try:
                return datetime.strptime(s, "%Y%m%d").date()
            except ValueError:
                return None
        # Excel 日期序列号（1900 日期系统，1900-01-01 序列号为 1）
        if 20000 <= n <= 80000:
            try:
                return (datetime(1899, 12, 30) + timedelta(days=n)).date()
            except (ValueError, OverflowError):
                return None
        return None
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return None
        cleaned = re.sub(r"[^\d]", "", s)
        if len(cleaned) >= 8 and cleaned[:8].isdigit():
            try:
                return datetime.strptime(cleaned[:8], "%Y%m%d").date()
            except ValueError:
                return None
        m = re.match(r"^(\d{4})[年/\-.](\d{1,2})[月/\-.](\d{1,2})日?$", s)
        if m:
            try:
                return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3))).date()
            except ValueError:
                return None
    return None


def _first_week_indicies(header_row):
    """表头行中识别前 3 个周出货量列（0-based 列索引）。

    按日期识别列，优先取相邻列间隔 7 天（周节奏）的最长连续段，取其前 3 列；
    无周节奏时回退最长连续日期段。兼容客户文件周列位置与格式变化（如 AD/AE/AF、
    字符串日期）。无日期列返回空。
    """
    cols = []  # (列号0based, 日期)
    for i, v in enumerate(header_row):
        d = _to_date(v)
        if d is not None:
            cols.append((i, d))
    if not cols:
        return []
    # 周节奏段
    weekly = None
    cur = []
    for item in cols:
        if cur and item[0] == cur[-1][0] + 1 and (item[1] - cur[-1][1]).days == 7:
            cur.append(item)
        else:
            if len(cur) > 1 and (weekly is None or len(cur) > len(weekly)):
                weekly = cur
            cur = [item]
    if len(cur) > 1 and (weekly is None or len(cur) > len(weekly)):
        weekly = cur
    if weekly:
        return [c for c, _ in weekly[:3]]
    # 回退：最长连续日期段
    longest, cur = [], []
    for item in cols:
        if cur and item[0] == cur[-1][0] + 1:
            cur.append(item)
        else:
            if len(cur) > len(longest):
                longest = cur
            cur = [item]
    if len(cur) > len(longest):
        longest = cur
    return [c for c, _ in longest[:3]]


def _read_template_sheets(excel_path, max_rows=200, max_cols=200, target_sheets=S2_SHEET_CANDIDATES, full=False, strict=False, filter_zero_weeks=False):
    """读取模板 sheet 内容，返回 (sheets, 错误消息)。

    sheets 结构: [{name, max_row, max_col, truncated, filtered_rows, rows}]
    优先只读目标 sheet（前端模板内容预览默认看排产2/SHEET1，预估出货量表看排产1/SHEET1）；
    找不到目标 sheet 时：
      - 文件仅一个 sheet → 无论叫什么（如 sheet1(2)）都视为目标 sheet；
      - 多个 sheet 且 strict=True → 返回错误（导入场景必须明确报错，避免静默展示旧数据）；
      - 多个 sheet 且 strict=False → 回退返回全部 sheet 以便排查。
    full=True 时读取全量行（不截断）。
    filter_zero_weeks=True 时按表头日期动态识别前 3 周列，过滤掉前 3 周均为 0/空 的数据行，
    并统计 filtered_rows（预估出货量导入场景使用，减少无效数据、加快展示）。
    单元格统一转字符串，日期格式化为 YYYY-MM-DD，避免前端展示溢出。
    """
    try:
        import openpyxl
    except ImportError:
        return None, "缺少 openpyxl 依赖，无法读取模板内容"
    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True, keep_links=False)
    except Exception as e:  # noqa: BLE001
        return None, f"模板读取失败: {e}"

    if full:
        max_rows = None

    # 优先定位目标 sheet（兼容名称中带空格等变体）；仅一个 sheet 时直接视为目标 sheet
    target = None
    norm_targets = {s.strip().replace(" ", "") for s in target_sheets}
    for name in wb.sheetnames:
        if name.strip().replace(" ", "") in norm_targets:
            target = name
            break
    if target is None and len(wb.sheetnames) == 1:
        target = wb.sheetnames[0]
    if target is None and strict:
        sheet_list = "、".join(wb.sheetnames)
        wb.close()
        return None, (f"未找到目标 sheet（{'/'.join(target_sheets)}），"
                      f"文件含多个 sheet（{sheet_list}），请检查文件")
    names = [target] if target else wb.sheetnames

    sheets = []
    try:
        for name in names:
            ws = wb[name]
            rows = []
            total_rows = 0
            total_cols = 0
            truncated = False
            filtered_rows = 0
            # 表头日期列：在 1~3 行中取日期单元格最多的行识别前 3 周列（兼容客户多行表头）
            # 注意：iter_rows 限定 max_col，避免读取超宽列（如 16365 列）导致极慢
            header_candidates = list(ws.iter_rows(values_only=True, max_row=3, max_col=max_cols))
            header_row = None
            best_n = 0
            for hrow in header_candidates:
                n = sum(1 for v in hrow if _to_date(v) is not None)
                if n > best_n:
                    best_n = n
                    header_row = list(hrow)
            week_cols = _first_week_indicies(header_row) if header_row else []
            # 表头有效列数：1~3 行中最后一个非空列。表头为空的尾部列（如 BL 列之后）不读取，
            # 避免导入/预览把空列一并展示（预估出货量表最后一列非空为止）。
            header_max_col = 0
            for hrow in header_candidates:
                for i, v in enumerate(hrow):
                    if v is not None and str(v).strip():
                        header_max_col = max(header_max_col, i + 1)
            effective_cols = min(max_cols, header_max_col) if header_max_col else max_cols

            # 动态识别库存相关列：差异量减排产量 / 预估减出货量减库存量 / 生产线量
            diff_col = estimate_col = line_col = None
            for hrow in header_candidates:
                for i, v in enumerate(hrow):
                    if v is None:
                        continue
                    s = str(v)
                    if diff_col is None and "差异量减排产量" in s:
                        diff_col = i
                    if estimate_col is None and "预估减出货量减库存量" in s:
                        estimate_col = i
                    if line_col is None and "生产线量" in s:
                        line_col = i
                if diff_col is not None and estimate_col is not None and line_col is not None:
                    break

            for row in ws.iter_rows(values_only=True, max_col=effective_cols):
                total_rows += 1
                total_cols = max(total_cols, effective_cols)
                # 数据行：前3周均为 0/空 → 过滤（保留表头行）
                if filter_zero_weeks and week_cols and total_rows > 1:
                    vals = [row[c] if c < len(row) else None for c in week_cols]
                    if all(v is None or (isinstance(v, str) and not str(v).strip()) or _num(v) == 0 for v in vals):
                        filtered_rows += 1
                        continue
                if max_rows is not None and len(rows) >= max_rows:
                    truncated = True
                    continue

                # 按客户规则重新计算「差异量减排产量」（仅数据行，不覆盖表头）：
                # 生产线量 <> #N/A 时，差异量减排产量 = 预估减出货量减库存量 - 生产线量；
                # 生产线量 = #N/A 时，差异量减排产量 = 预估减出货量减库存量。
                row_list = list(row)
                if total_rows > 1 and diff_col is not None and estimate_col is not None and line_col is not None \
                        and diff_col < len(row_list) and estimate_col < len(row_list) and line_col < len(row_list):
                    line_val = _to_number(row_list[line_col])
                    estimate_val = _to_number(row_list[estimate_col])
                    if not _is_na(line_val) and line_val is not None:
                        computed = (estimate_val - line_val) if estimate_val is not None else line_val
                    else:
                        computed = estimate_val
                    row_list[diff_col] = _fmt_number(computed)
                cells = []
                for v in row_list:
                    if v is None:
                        cells.append("")
                    elif isinstance(v, datetime):
                        cells.append(v.strftime("%Y-%m-%d"))
                    else:
                        cells.append(str(v))
                rows.append(cells)
            sheets.append({
                "name": name,
                "max_row": total_rows,
                "max_col": total_cols,
                "truncated": truncated,
                "filtered_rows": filtered_rows,
                "week_cols": [c + 1 for c in week_cols],  # 1-based，供前端提示识别结果
                "rows": rows,
            })
    finally:
        wb.close()
    return sheets, None


def _config_path(name):
    """配置文件名 → 绝对路径。"""
    return os.path.join(_scheduling_config_dir(), CONFIG_FILES[name][0])


def _read_config(name, default=None):
    """读取配置文件，不存在返回 default。"""
    p = _config_path(name)
    if not os.path.isfile(p):
        return default
    try:
        with open(p, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:  # noqa: BLE001
        logger.warning(f"[AirbagScheduling] 读取配置 {name} 失败: {e}")
        return default


def _html_escape(value):
    """HTML 转义，防止物料号/备注中的特殊字符破坏页面。"""
    return (str(value or "")
            .replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace('"', "&quot;").replace("'", "&#39;"))


def _num(value, default=0):
    """安全转数值，非数字返回 default。"""
    if value is None or value == "":
        return default
    try:
        return float(str(value).replace(",", "").strip())
    except (ValueError, TypeError):
        return default


def _is_na(value):
    """判断单元格值是否为无效值（None、空串、#N/A）。"""
    if value is None:
        return True
    if isinstance(value, str):
        s = value.strip()
        return not s or s.upper() == "#N/A"
    return False


def _to_number(value):
    """尝试转数值；无效（None/#N/A/空串/非数字）返回 None。"""
    if _is_na(value):
        return None
    if isinstance(value, (int, float)):
        return value
    s = str(value).replace(",", "").strip()
    if s.upper() == "#N/A":
        return None
    try:
        return float(s)
    except (ValueError, TypeError):
        return None


def _fmt_number(value):
    """数值格式化为展示字符串；None/非数值返回空串。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _build_report_html(payload):
    """根据引擎 run 输出生成可视化排产分析报告 HTML（自包含单文件）。

    内容：① 排产摘要 KPI  ② 产线使用分析（已排/空闲）  ③ 班组负荷与人天
          ④ 物料排产明细  ⑤ 提示与建议
    """
    results = payload.get("results") or []
    skipped = payload.get("skipped") or []
    missing = payload.get("missing_capacity") or []
    warnings = payload.get("warnings") or []
    window = payload.get("window") or {}
    window_days = window.get("days") or 0
    window_text = f"{_html_escape(window.get('start', '-'))} ~ {_html_escape(window.get('end', '-'))}（{window_days}天）"

    teams_cfg = _read_config("teams", []) or []
    team_count = {str(t.get("team", "")).strip(): t.get("count") for t in teams_cfg if t.get("team")}

    # ---- 统计：产线 / 班组 / 汇总 ----
    line_stat = {}   # line -> {days, night_days, teams:set, saps:set}
    team_stat = {}   # team -> {days}
    total_days = 0
    stockup_cnt = shebian_cnt = night_cnt = 0
    for r in results:
        if r.get("stockup"):
            stockup_cnt += 1
        if r.get("shebian"):
            shebian_cnt += 1
        if r.get("trigger_night"):
            night_cnt += 1
        for l in r.get("lines", []):
            days = int(l.get("early_days") or 0) + int(l.get("night_days") or 0)
            if days <= 0:
                continue
            total_days += days
            ls = line_stat.setdefault(str(l.get("line") or "?"),
                                      {"days": 0, "night_days": 0, "teams": set(), "saps": set()})
            ls["days"] += days
            ls["night_days"] += int(l.get("night_days") or 0)
            ls["saps"].add(r.get("sap", "?"))
            for t in (l.get("team") or []):
                t = str(t)
                ls["teams"].add(t)
                ts = team_stat.setdefault(t, {"days": 0, "count": None})
                ts["days"] += days
                if ts["count"] is None and t in team_count:
                    ts["count"] = team_count[t]

    # 估算总人天：线天 × 班组人数（人数缺失按 1 人）
    total_man_days = 0
    for t, ts in team_stat.items():
        total_man_days += ts["days"] * int(ts["count"] or team_count.get(t) or 1)

    # ---- 渲染辅助 ----
    def bar(ratio, max_ratio=None):
        """条形图 HTML；ratio 0~1，max_ratio 用于归一化。"""
        mr = max_ratio or max(ratio, 0.001)
        w = max(2.0, min(100.0, ratio / mr * 100))
        cls = "high" if ratio >= 0.9 else ("mid" if ratio >= 0.6 else "")
        return (f'<div class="bar-bg"><div class="bar-fill {cls}" style="width:{w:.1f}%">'
                f'</div><div class="bar-text">{ratio * 100:.0f}%</div></div>')

    def badge_ratio(ratio):
        if ratio >= 0.9:
            return '<span class="badge danger">高负荷</span>'
        if ratio >= 0.6:
            return '<span class="badge warn">中负荷</span>'
        if ratio <= 0.0:
            return '<span class="badge lv0">空闲</span>'
        return '<span class="badge ok">正常</span>'

    # ---- 产线使用 ----
    line_rows = []
    idle_lines = []
    max_line_ratio = 0.001
    for line, ls in sorted(line_stat.items()):
        ratio = ls["days"] / window_days if window_days else 0
        max_line_ratio = max(max_line_ratio, ratio)
    for line, ls in sorted(line_stat.items(), key=lambda kv: -kv[1]["days"]):
        ratio = ls["days"] / window_days if window_days else 0
        teams_txt = "、".join(sorted(ls["teams"])) or "-"
        saps_txt = "、".join(sorted(ls["saps"])) or "-"
        if ratio <= 0:
            idle_lines.append(line)
        line_rows.append(
            f'<tr><td><b>{_html_escape(line)}</b></td>'
            f'<td>{ls["days"]}（夜{ls["night_days"]}）</td>'
            f'<td>{bar(ratio, max_line_ratio)}</td>'
            f'<td>{badge_ratio(ratio)}</td>'
            f'<td>{len(ls["saps"])}</td>'
            f'<td style="color:#64748b;font-size:12px;">{teams_txt}</td>'
            f'<td style="color:#64748b;font-size:12px;">{saps_txt}</td></tr>')
    # 全产线清单（含空闲）：产能配置中的产线
    capacity_cfg = _read_config("capacity", {}) or {}
    all_lines = set()
    for _sap, arr in capacity_cfg.items():
        for it in arr or []:
            if it.get("line"):
                all_lines.add(str(it["line"]))
    idle_extra = sorted(all_lines - set(line_stat.keys()))
    if idle_extra:
        for line in idle_extra:
            line_rows.append(
                f'<tr><td><b>{_html_escape(line)}</b></td>'
                f'<td>0</td><td>{bar(0)}</td>'
                f'<td><span class="badge lv0">空闲</span></td>'
                f'<td>0</td><td>-</td><td>-</td></tr>')
        idle_lines += idle_extra

    # ---- 班组负荷 ----
    team_rows = []
    max_team_ratio = 0.001
    for t, ts in team_stat.items():
        ratio = ts["days"] / total_days if total_days else 0
        max_team_ratio = max(max_team_ratio, ratio)
    for t, ts in sorted(team_stat.items(), key=lambda kv: -kv[1]["days"]):
        ratio = ts["days"] / total_days if total_days else 0
        cnt = int(ts["count"] or team_count.get(t) or 1)
        md = ts["days"] * cnt
        team_rows.append(
            f'<tr><td><b>{_html_escape(t)}</b></td>'
            f'<td>{ts["days"]} 线天</td>'
            f'<td>{cnt} 人</td>'
            f'<td>{bar(ratio, max_team_ratio)}</td>'
            f'<td>{md} 人天</td></tr>')

    # ---- 物料明细 ----
    detail_rows = []
    for r in results:
        badges = []
        if r.get("stockup"):
            badges.append('<span class="badge ok">备货</span>')
        if r.get("shebian"):
            badges.append('<span class="badge warn">设变</span>')
        if r.get("trigger_night"):
            badges.append('<span class="badge lv1">夜班</span>')
        lines_txt = "；".join(
            f'{_html_escape(l.get("line", "?"))}（早{l.get("early_days", 0)}天+夜{l.get("night_days", 0)}天'
            f'，班组{"+".join(l.get("team") or []) or "-"}）'
            for l in r.get("lines", [])
        ) or "-"
        detail_rows.append(
            f'<tr><td><b>{_html_escape(r.get("sap", "?"))}</b> {"".join(badges)}</td>'
            f'<td>{_html_escape(r.get("g", "")) or "-"}</td>'
            f'<td>{_html_escape(r.get("o", "")) or "-"}</td>'
            f'<td>{_num(r.get("ay"))}</td>'
            f'<td>{_num(r.get("w1"))} / {_num(r.get("w2"))}</td>'
            f'<td>{_num(r.get("need"))}</td>'
            f'<td style="font-size:12px;">{lines_txt}</td>'
            f'<td style="color:#64748b;font-size:12px;">{_html_escape(r.get("reason", ""))}</td></tr>')

    # ---- 建议 ----
    notes = []
    if idle_lines:
        notes.append(f'<div class="alert-box med"><div class="title">🟢 空闲产线（{len(idle_lines)} 条）</div>'
                     f'<div class="sug">{"、".join(_html_escape(x) for x in sorted(idle_lines))} 未排产，'
                     f'可承接急单或安排维保。</div></div>')
    busy = [line for line, ls in line_stat.items()
            if window_days and ls["days"] / window_days >= 0.9]
    if busy:
        notes.append(f'<div class="alert-box high"><div class="title">🔴 高负荷产线</div>'
                     f'<div class="sug">{"、".join(_html_escape(x) for x in sorted(busy))} 占用率 ≥90%，'
                     f'接近满产，注意产能瓶颈。</div></div>')
    if skipped:
        skipped_txt = "；".join(
            f"{_html_escape(s.get('sap', '?'))}：{_html_escape(s.get('reason', ''))}" for s in skipped)
        notes.append(f'<div class="alert-box med"><div class="title">⚠️ 跳过需求（{len(skipped)}）</div>'
                     f'<div class="sug">{skipped_txt}</div></div>')
    if missing:
        notes.append(f'<div class="alert-box med"><div class="title">⚠️ 缺失产能配置（{len(missing)}）</div>'
                     f'<div class="sug">{"、".join(_html_escape(x) for x in missing)} 无产能配置，需在「产能配置」中维护后重排。</div></div>')
    if warnings:
        notes.append(f'<div class="alert-box med"><div class="title">ℹ️ 引擎提示</div>'
                     f'<div class="sug">{"；".join(_html_escape(x) for x in warnings[:5])}</div></div>')
    if not notes:
        notes.append('<div class="alert-box" style="border-color:#10b981;background:#f0fdf4;">'
                     '<div class="title">✅ 排产正常</div>'
                     '<div class="sug">全部需求已排产，无跳过、无缺失产能。</div></div>')

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    return f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="UTF-8">
<title>气袋排产分析报告</title>
<style>
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
body {{ font-family: -apple-system, "Segoe UI", "Microsoft YaHei", sans-serif; background: #f5f7fa; color: #1e293b; line-height: 1.6; }}
.container {{ max-width: 1280px; margin: 0 auto; padding: 24px; }}
header {{ background: linear-gradient(135deg, #0d9488 0%, #14b8a6 100%); color: white; padding: 32px; border-radius: 12px; margin-bottom: 24px; box-shadow: 0 4px 12px rgba(0,0,0,0.08); }}
header h1 {{ font-size: 26px; margin-bottom: 8px; }}
header .meta {{ opacity: 0.92; font-size: 14px; }}
header .meta span {{ margin-right: 16px; }}
section {{ background: white; padding: 24px; border-radius: 12px; margin-bottom: 20px; box-shadow: 0 1px 3px rgba(0,0,0,0.05); }}
h2 {{ font-size: 19px; margin-bottom: 16px; color: #1e293b; border-left: 4px solid #0d9488; padding-left: 12px; }}
.kpi-grid {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(170px, 1fr)); gap: 16px; }}
.kpi {{ background: #f8fafc; padding: 16px; border-radius: 8px; border-left: 4px solid #0d9488; }}
.kpi.warn {{ border-left-color: #f59e0b; background: #fffbeb; }}
.kpi.danger {{ border-left-color: #ef4444; background: #fef2f2; }}
.kpi.ok {{ border-left-color: #10b981; background: #f0fdf4; }}
.kpi .label {{ font-size: 13px; color: #64748b; margin-bottom: 6px; }}
.kpi .value {{ font-size: 25px; font-weight: 700; color: #1e293b; }}
.kpi .sub {{ font-size: 12px; color: #94a3b8; margin-top: 4px; }}
table {{ width: 100%; border-collapse: collapse; font-size: 13px; }}
th {{ background: #f1f5f9; padding: 10px 12px; text-align: left; font-weight: 600; color: #475569; border-bottom: 2px solid #e2e8f0; }}
td {{ padding: 10px 12px; border-bottom: 1px solid #f1f5f9; }}
tr:hover {{ background: #f8fafc; }}
.badge {{ display: inline-block; padding: 2px 8px; border-radius: 4px; font-size: 12px; font-weight: 500; }}
.badge.ok {{ background: #dcfce7; color: #166534; }}
.badge.warn {{ background: #fef3c7; color: #92400e; }}
.badge.danger {{ background: #fee2e2; color: #991b1b; }}
.badge.lv0 {{ background: #e0e7ff; color: #3730a3; }}
.badge.lv1 {{ background: #dbeafe; color: #1e40af; }}
.bar-bg {{ background: #e2e8f0; height: 18px; border-radius: 4px; overflow: hidden; position: relative; }}
.bar-fill {{ background: linear-gradient(90deg, #0d9488, #2dd4bf); height: 100%; }}
.bar-fill.high {{ background: linear-gradient(90deg, #ef4444, #f87171); }}
.bar-fill.mid {{ background: linear-gradient(90deg, #f59e0b, #fbbf24); }}
.bar-text {{ position: absolute; top: 50%; left: 8px; transform: translateY(-50%); color: #1e293b; font-size: 12px; font-weight: 500; }}
.alert-box {{ padding: 12px 16px; border-radius: 8px; margin-bottom: 8px; border-left: 4px solid; }}
.alert-box.high {{ background: #fef2f2; border-color: #ef4444; }}
.alert-box.med {{ background: #fffbeb; border-color: #f59e0b; }}
.alert-box .title {{ font-weight: 600; margin-bottom: 4px; }}
.alert-box .sug {{ font-size: 13px; color: #64748b; }}
.note {{ background: #f0fdf4; border-left: 4px solid #10b981; padding: 12px 16px; border-radius: 4px; margin: 12px 0; font-size: 14px; color: #14532d; }}
@media print {{ header {{ -webkit-print-color-adjust: exact; }} body {{ background: white; }} }}
</style></head>
<body><div class="container">
<header>
  <h1>📊 气袋排产分析报告</h1>
  <div class="meta"><span>🕐 生成时间: {now}</span><span>📅 排产窗口: {window_text}</span>
  <span>📦 已排物料 {len(results)} 个 · 线天 {total_days} · 人天 {total_man_days}</span></div>
</header>

<section><h2>① 排产摘要</h2><div class="kpi-grid">
  <div class="kpi"><div class="label">排产物料数</div><div class="value">{len(results)}</div><div class="sub">已生成排产计划</div></div>
  <div class="kpi"><div class="label">总排线天</div><div class="value">{total_days}</div><div class="sub">早班+夜班合计</div></div>
  <div class="kpi ok"><div class="label">估算总人天</div><div class="value">{total_man_days}</div><div class="sub">按班组人数折算</div></div>
  <div class="kpi"><div class="label">排产窗口</div><div class="value">{window_days}</div><div class="sub">天</div></div>
  <div class="kpi {'' if not skipped else 'warn'}"><div class="label">跳过需求</div><div class="value">{len(skipped)}</div><div class="sub">{'; '.join(_html_escape(s.get('sap','')) for s in skipped[:4]) or '-'}</div></div>
  <div class="kpi {'' if not missing else 'danger'}"><div class="label">缺失产能</div><div class="value">{len(missing)}</div><div class="sub">{'; '.join(_html_escape(x) for x in missing[:4]) or '-'}</div></div>
  <div class="kpi"><div class="label">备货物料</div><div class="value">{stockup_cnt}</div><div class="sub">提前排产</div></div>
  <div class="kpi"><div class="label">设变物料</div><div class="value">{shebian_cnt}</div><div class="sub">需重点关注</div></div>
  <div class="kpi"><div class="label">夜班物料</div><div class="value">{night_cnt}</div><div class="sub">触发夜班</div></div>
</div></section>

<section><h2>② 产线使用分析</h2>
  <p style="color:#64748b;font-size:13px;margin-bottom:12px;">共 {len(line_stat) + len(idle_extra)} 条产线参与评估；条柱为排产天数占窗口比例（已归一化）。</p>
  <table><thead><tr><th>产线</th><th>排产天数</th><th>占用</th><th>状态</th><th>物料数</th><th>班组</th><th>SAP物料</th></tr></thead>
  <tbody>{''.join(line_rows)}</tbody></table>
</section>

<section><h2>③ 班组负荷与人天</h2>
  <p style="color:#64748b;font-size:13px;margin-bottom:12px;">人天 = 排产线天 × 班组人数；人数来自「班组产线」配置，缺失按 1 人估算。</p>
  <table><thead><tr><th>班组</th><th>排产线天</th><th>人数</th><th>负荷占比</th><th>估算人天</th></tr></thead>
  <tbody>{''.join(team_rows)}</tbody></table>
  <div class="note">💡 总资源计划：窗口 {window_days} 天，共需投入约 <b>{total_man_days} 人天</b>，折算日均 <b>{round(total_man_days / window_days, 1) if window_days else 0} 人</b>。</div>
</section>

<section><h2>④ 物料排产明细</h2>
  <table><thead><tr><th>SAP物料</th><th>备货(G)</th><th>备注(O)</th><th>可用库存</th><th>W1/W2</th><th>需排产</th><th>产线安排</th><th>原因</th></tr></thead>
  <tbody>{''.join(detail_rows)}</tbody></table>
</section>

<section><h2>⑤ 提示与建议</h2>
  {''.join(notes)}
</section>

</div></body></html>"""


def _write_config(name, data):
    """写入配置文件。"""
    p = _config_path(name)
    with open(p, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return p


def _run_engine(mode, excel_path, output_path=None, material=None, demands_path=None, config_only=False):
    """调用技能引擎脚本。

    返回 (payload dict, 错误消息)。
    config_only=True 时以 --config-only 调用，仅提取配置（不要求排产2 结构/需求）。
    """
    skill_dir = _resolve_skill_dir()
    if not skill_dir:
        return None, "未找到技能目录 pmc-scheduler-hmt-qd"
    script = os.path.join(skill_dir, "scripts", "schedule_production.py")
    if not os.path.isfile(script):
        return None, f"未找到引擎脚本: {script}"

    cmd = [sys.executable, script, mode, "--excel", excel_path, "--json"]
    if config_only:
        cmd.append("--config-only")
    else:
        # 仅传入已存在的配置文件；不存在时引擎回退读取 Excel sheet / 内置默认
        for name, (_, flag) in CONFIG_FILES.items():
            p = _config_path(name)
            if os.path.isfile(p):
                cmd += [flag, p]
        if demands_path and os.path.isfile(demands_path):
            cmd += ["--demands", demands_path]
    if output_path:
        cmd += ["--output", output_path]
    if material:
        cmd += ["--material", material]

    logger.info(f"[AirbagScheduling] 调用引擎: {' '.join(cmd)}")
    try:
        # PYTHONIOENCODING 强制引擎 stdout/stderr 使用 UTF-8，避免 Windows GBK 中文乱码
        env = dict(os.environ)
        env["PYTHONIOENCODING"] = "utf-8"
        proc = subprocess.run(cmd, capture_output=True, encoding="utf-8", timeout=300, env=env)
    except subprocess.TimeoutExpired:
        return None, "引擎执行超时（300 秒）"
    except Exception as e:  # noqa: BLE001
        logger.error(f"[AirbagScheduling] 引擎执行异常: {e}")
        return None, f"引擎执行异常: {e}"

    stderr_tail = (proc.stderr or "").strip().splitlines()[-5:]
    if proc.returncode != 0:
        detail = "\n".join(stderr_tail) if stderr_tail else proc.stdout[-500:]
        return None, f"引擎执行失败（exit={proc.returncode}）:\n{detail}"

    try:
        payload = json.loads(proc.stdout)
    except json.JSONDecodeError:
        return None, f"引擎输出不是合法 JSON:\n{proc.stdout[-800:]}"
    if stderr_tail:
        payload["engine_warnings"] = stderr_tail
    return payload, None


# ---------------------------------------------------------------------------
# Handler
# ---------------------------------------------------------------------------
class AirbagSchedulingHandler:
    """气袋排产 Handler：GET/PUT/POST 按 action 分发。"""

    def _check_scene_permission(self):
        _require_auth()
        _require_permission(SCENE_PERMISSION)

    def _check_config_permission(self):
        _require_auth()
        _require_permission(CONFIG_MANAGE_PERMISSION)

    # ---------------------------------------------------------- GET
    def GET(self, action=""):
        self._check_scene_permission()
        web.header("Content-Type", "application/json; charset=utf-8")

        try:
            if action in ("", "config"):
                rules = _read_config("rules", dict(DEFAULT_RULES))
                return json.dumps({"status": "success", "data": rules}, ensure_ascii=False)
            if action == "capacity":
                return json.dumps({"status": "success", "data": _read_config("capacity", [])}, ensure_ascii=False)
            if action == "teams":
                return json.dumps({"status": "success", "data": _read_config("teams", [])}, ensure_ascii=False)
            if action == "merge-rules":
                return json.dumps({"status": "success", "data": _read_config("merge_rules", [])}, ensure_ascii=False)
            if action == "history":
                history = self._read_history()
                return json.dumps({"status": "success", "data": history}, ensure_ascii=False)
            if action == "template":
                return json.dumps({"status": "success", "data": _template_info()}, ensure_ascii=False)
            if action == "demands":
                return self._do_demands_info()
            return json.dumps({"status": "error", "message": f"未知操作: {action}"}, ensure_ascii=False)
        except web.HTTPError:
            raise
        except Exception as e:  # noqa: BLE001
            logger.error(f"[AirbagScheduling] GET {action} 失败: {e}")
            return json.dumps({"status": "error", "message": f"读取失败: {e}"}, ensure_ascii=False)

    # ---------------------------------------------------------- PUT
    def PUT(self, action=""):
        self._check_config_permission()
        web.header("Content-Type", "application/json; charset=utf-8")

        try:
            body = json.loads(web.data() or b"{}")
            data = body.get("data")
            if data is None:
                return json.dumps({"status": "error", "message": "缺少 data 字段"}, ensure_ascii=False)

            name_map = {
                "": "rules",
                "config": "rules",
                "capacity": "capacity",
                "teams": "teams",
                "merge-rules": "merge_rules",
            }
            name = name_map.get(action)
            if not name:
                return json.dumps({"status": "error", "message": f"未知操作: {action}"}, ensure_ascii=False)
            _write_config(name, data)
            return json.dumps({"status": "success", "message": "配置已保存"}, ensure_ascii=False)
        except web.HTTPError:
            raise
        except Exception as e:  # noqa: BLE001
            logger.error(f"[AirbagScheduling] PUT {action} 失败: {e}")
            return json.dumps({"status": "error", "message": f"保存失败: {e}"}, ensure_ascii=False)

    # ---------------------------------------------------------- POST
    def POST(self, action=""):
        web.header("Content-Type", "application/json; charset=utf-8")

        try:
            if action == "parse":
                return self._do_parse()
            if action == "template-content":
                return self._do_template_content()
            if action == "demands":
                return self._do_demands_upload()
            if action == "sync-from-excel":
                return self._do_sync_from_excel()
            if action == "run":
                return self._do_run()
            if action == "template":
                return self._do_upload_template()
            if action == "template/reset":
                return self._do_reset_template()
            return json.dumps({"status": "error", "message": f"未知操作: {action}"}, ensure_ascii=False)
        except web.HTTPError:
            raise
        except Exception as e:  # noqa: BLE001
            logger.error(f"[AirbagScheduling] POST {action} 失败: {e}")
            return json.dumps({"status": "error", "message": f"操作失败: {e}"}, ensure_ascii=False)

    # ---------------------------------------------------------- 各动作
    def _do_parse(self):
        """解析预览：可选上传新 Excel，否则使用当前默认模板。"""
        self._check_scene_permission()
        body = json.loads(web.data() or b"{}")
        excel_path, err = _save_uploaded_file(body)
        if err:
            return json.dumps({"status": "error", "message": err}, ensure_ascii=False)
        if not excel_path:
            excel_path, source = _resolve_template()
            if not excel_path:
                return json.dumps({"status": "error", "message": "未找到可用模板，请先上传排产模板"}, ensure_ascii=False)

        payload, err = _run_engine("parse", excel_path, material=body.get("material") or None, demands_path=_demands_path())
        if err:
            return json.dumps({"status": "error", "message": err}, ensure_ascii=False)
        payload["template_source"] = "uploaded_now" if body.get("files") else _template_info()["source"]
        return json.dumps({"status": "success", "data": payload}, ensure_ascii=False)

    def _do_template_content(self):
        """模板内容预览：读取模板各 sheet 原始内容（截断），供工作台展示。"""
        self._check_scene_permission()
        body = json.loads(web.data() or b"{}")
        excel_path, err = _save_uploaded_file(body)
        if err:
            return json.dumps({"status": "error", "message": err}, ensure_ascii=False)
        if not excel_path:
            excel_path, source = _resolve_template()
            if not excel_path:
                return json.dumps({"status": "error", "message": "未找到可用模板，请先上传排产模板"}, ensure_ascii=False)

        sheets, err = _read_template_sheets(excel_path, max_cols=S2_PREVIEW_MAX_COLS)
        if err:
            return json.dumps({"status": "error", "message": err}, ensure_ascii=False)
        info = _template_info()
        return json.dumps(
            {
                "status": "success",
                "data": {
                    "template_source": "uploaded_now" if body.get("files") else info["source"],
                    "sheets": sheets,
                },
            },
            ensure_ascii=False,
        )

    def _do_demands_info(self):
        """预估出货量表信息 + 排产1 内容预览（场景权限）。

        支持 GET ?full=1 返回全量行（默认截断 200 行）；
        读取时按表头日期过滤前 3 周均无出货的记录（filtered_rows）。
        """
        info = _demands_info()
        sheets = []
        p = _demands_path()
        if p:
            full = str(web.input(full="").get("full", "")).lower() in ("1", "true")
            sheets, err = _read_template_sheets(
                p, target_sheets=DEMAND_SHEET_CANDIDATES, full=full, strict=True, filter_zero_weeks=True)
            if err:
                return json.dumps({"status": "error", "message": err}, ensure_ascii=False)
        return json.dumps(
            {"status": "success", "data": dict(info, sheets=sheets)},
            ensure_ascii=False,
        )

    def _do_demands_upload(self):
        """上传预估出货量表：保存为 demands.xlsx，作为排产1 数据来源（场景权限）。

        导入时按表头日期识别前 3 周列，过滤掉前 3 周均为 0/空 的记录，并统计条数反馈前端。
        """
        self._check_scene_permission()
        body = json.loads(web.data() or b"{}")
        excel_path, err = _save_uploaded_file(body)
        if err:
            return json.dumps({"status": "error", "message": err}, ensure_ascii=False)
        if not excel_path:
            return json.dumps({"status": "error", "message": "请上传预估出货量表（含排产1/SHEET1 sheet 的 Excel）"}, ensure_ascii=False)

        import shutil
        dest = os.path.join(_scheduling_config_dir(), "demands.xlsx")
        shutil.copy2(excel_path, dest)
        # 去掉 SAP 物料号（B 列）前导零，界面展示更简洁（如 000000005000001256 → 5000001256）
        _normalize_sap_zeros(dest)
        full = body.get("full") in (True, 1, "1", "true")
        sheets, err = _read_template_sheets(
            dest, target_sheets=DEMAND_SHEET_CANDIDATES, full=full, strict=True, filter_zero_weeks=True)
        if err:
            return json.dumps({"status": "error", "message": err}, ensure_ascii=False)
        # 条数统计：共 N 条 = 总行-表头；过滤 K 条 = 前3周无出货；有效 M 条
        sheet = sheets[0] if sheets else {}
        total = max(0, int(sheet.get("max_row", 0)) - 1)
        filtered = int(sheet.get("filtered_rows", 0))
        valid = total - filtered
        message = (f"预估出货量表导入成功：共 {total} 条记录，"
                   f"其中前3周无出货已过滤 {filtered} 条，有效需求 {valid} 条")
        return json.dumps(
            {
                "status": "success",
                "message": message,
                "data": dict(_demands_info(), sheets=sheets, total_rows=total, filtered_rows=filtered, valid_rows=valid),
            },
            ensure_ascii=False,
        )

    def _do_sync_from_excel(self):
        """从上传 Excel 同步产能/班组/拼线配置到系统配置表。

        支持独立配置表（如仅「班组产线」一个 sheet 的文件）：
        以引擎 --config-only 提取配置，不要求完整排产模板结构。
        body.target 指定同步目标：capacity / teams / merge（缺省兼容旧前端=全部）。
        """
        self._check_config_permission()
        body = json.loads(web.data() or b"{}")
        target = body.get("target") or ""
        excel_path, err = _save_uploaded_file(body)
        if err:
            return json.dumps({"status": "error", "message": err}, ensure_ascii=False)
        if not excel_path:
            excel_path, source = _resolve_template()
            if not excel_path:
                return json.dumps({"status": "error", "message": "未找到可用模板，请先上传排产模板"}, ensure_ascii=False)

        payload, err = _run_engine("parse", excel_path, config_only=True)
        if err:
            return json.dumps({"status": "error", "message": err}, ensure_ascii=False)

        # 按目标保存；merge 派生自班组备注，随班组一起同步
        if target == "capacity":
            keys = ("capacity",)
        elif target == "teams":
            keys = ("teams", "merge_rules")
        elif target == "merge":
            keys = ("merge_rules",)
        else:
            keys = ("capacity", "teams", "merge_rules")

        saved = {}
        labels = {"capacity": "产能", "teams": "班组", "merge_rules": "拼线"}
        for key in keys:
            val = payload.get(key)
            if val is None:
                continue
            if isinstance(val, dict) and not val:
                continue
            if isinstance(val, list) and not val:
                continue
            _write_config(key, val)
            if isinstance(val, list):
                saved[key] = len(val)
            elif isinstance(val, dict):
                saved[key] = len(val)  # 产能为 {sap: [...]}，按物料数统计
            else:
                saved[key] = 1
        if not saved:
            return json.dumps(
                {"status": "error", "message": "文件中未识别到可同步的配置数据（请确认含「产能」或「班组产线」sheet）"},
                ensure_ascii=False,
            )
        detail = "、".join(f"{labels.get(k, k)} {v} 条" for k, v in saved.items())
        return json.dumps(
            {"status": "success", "message": f"已从 Excel 同步：{detail}", "saved": saved},
            ensure_ascii=False,
        )

    def _do_run(self):
        """执行排产：生成新排产表 + 结果 JSON + 写入历史。"""
        self._check_scene_permission()
        body = json.loads(web.data() or b"{}")
        excel_path, err = _save_uploaded_file(body)
        if err:
            return json.dumps({"status": "error", "message": err}, ensure_ascii=False)
        if not excel_path:
            excel_path, source = _resolve_template()
            if not excel_path:
                return json.dumps({"status": "error", "message": "未找到可用模板，请先上传排产模板"}, ensure_ascii=False)

        output = os.path.join(_result_dir(), f"排产结果_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        payload, err = _run_engine(
            "run", excel_path, output_path=output, material=body.get("material") or None,
            demands_path=_demands_path(),
        )
        if err:
            return json.dumps({"status": "error", "message": err}, ensure_ascii=False)

        payload["output_path"] = output
        payload["template_source"] = _template_info()["source"]
        # 最新排产2 表内容：前端在模板内容区展示执行生成的结果
        payload["result_sheets"], sheets_err = _read_template_sheets(output, target_sheets=S2_SHEET_CANDIDATES, max_cols=S2_PREVIEW_MAX_COLS, strict=True)
        if sheets_err:
            logger.warning(f"[AirbagScheduling] 读取排产结果排产2 表失败: {sheets_err}")
        # 生成可视化排产分析报告（HTML）
        try:
            report_path = os.path.join(
                _result_dir(), f"排产分析报告_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html")
            with open(report_path, "w", encoding="utf-8") as f:
                f.write(_build_report_html(payload))
            payload["report_path"] = report_path
            payload["report_name"] = os.path.basename(report_path)
        except Exception as e:  # noqa: BLE001
            logger.warning(f"[AirbagScheduling] 生成排产分析报告失败: {e}")
        # 排产明细 JSON：完整结构化结果存文件，回传聊天窗只给路径（避免正文过长）
        try:
            detail_path = os.path.join(
                _result_dir(), f"排产明细_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json")
            with open(detail_path, "w", encoding="utf-8") as f:
                json.dump({
                    "window": payload.get("window"),
                    "written_rows": payload.get("written_rows"),
                    "results": payload.get("results", []),
                    "skipped": payload.get("skipped", []),
                    "warnings": payload.get("warnings", []),
                }, f, ensure_ascii=False, indent=2)
            payload["detail_path"] = detail_path
            payload["detail_name"] = os.path.basename(detail_path)
        except Exception as e:  # noqa: BLE001
            logger.warning(f"[AirbagScheduling] 生成排产明细 JSON 失败: {e}")
        self._append_history(payload, output)
        return json.dumps({"status": "success", "data": payload}, ensure_ascii=False)

    def _do_upload_template(self):
        """更换模板：保存上传文件为默认模板。"""
        self._check_config_permission()
        body = json.loads(web.data() or b"{}")
        excel_path, err = _save_uploaded_file(body)
        if err:
            return json.dumps({"status": "error", "message": err}, ensure_ascii=False)
        if not excel_path:
            return json.dumps({"status": "error", "message": "请上传新的排产模板文件"}, ensure_ascii=False)

        import shutil
        dest = os.path.join(_scheduling_config_dir(), "template.xlsx")
        shutil.copy2(excel_path, dest)
        return json.dumps(
            {"status": "success", "message": "模板已更新为默认模板", "data": _template_info()},
            ensure_ascii=False,
        )

    def _do_reset_template(self):
        """恢复内置模板：删除用户上传的模板副本。"""
        self._check_config_permission()
        uploaded = os.path.join(_scheduling_config_dir(), "template.xlsx")
        if os.path.isfile(uploaded):
            os.remove(uploaded)
        return json.dumps(
            {"status": "success", "message": "已恢复内置模板", "data": _template_info()},
            ensure_ascii=False,
        )

    # ---------------------------------------------------------- 历史记录
    def _match_report_path(self, entry):
        """旧记录补齐 report_path：按输出文件目录 + 时间就近匹配「排产分析报告_*.html」。"""
        out = entry.get("output_path")
        if not out:
            return None
        d = os.path.dirname(out)
        if not os.path.isdir(d):
            return None
        try:
            t = datetime.strptime(entry.get("time", ""), "%Y-%m-%d %H:%M:%S")
        except (ValueError, TypeError):
            return None
        best, best_gap = None, 3600  # 匹配窗口：±1 小时内最近的一个
        try:
            for f in os.listdir(d):
                if f.startswith("排产分析报告_") and f.endswith(".html"):
                    fp = os.path.join(d, f)
                    gap = abs((datetime.fromtimestamp(os.path.getmtime(fp)) - t).total_seconds())
                    if gap < best_gap:
                        best_gap = gap
                        best = fp
        except OSError:
            return None
        return best

    def _read_history(self):
        p = os.path.join(_result_dir(), "history.json")
        if not os.path.isfile(p):
            return []
        try:
            with open(p, "r", encoding="utf-8") as f:
                history = json.load(f)
        except Exception:  # noqa: BLE001
            return []
        # 旧记录补齐 report_path（新记录已由 _append_history 写入）；补齐后持久化一次
        changed = False
        for entry in history:
            if not entry.get("report_path"):
                rp = self._match_report_path(entry)
                if rp:
                    entry["report_path"] = rp
                    changed = True
        if changed:
            try:
                with open(p, "w", encoding="utf-8") as f:
                    json.dump(history, f, ensure_ascii=False, indent=2)
            except Exception:  # noqa: BLE001
                pass
        return history

    def _append_history(self, payload, output_path):
        """写入历史记录：时间、参数快照、结果摘要（保留最近 50 条）。"""
        history = self._read_history()
        entry = {
            "id": datetime.now().strftime("%Y%m%d%H%M%S"),
            "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "output_path": output_path,
            "report_path": payload.get("report_path"),  # 排产分析报告 HTML（供下载链接）
            "window": payload.get("window"),
            "rules": _read_config("rules"),
            # parse 用 demands，run 用 results
            "demand_count": len(payload.get("demands", payload.get("results", []))),
            "skipped_count": len(payload.get("skipped", [])),
            "warning_count": len(payload.get("warnings", [])),
            "summary": [
                {"sap": d.get("sap"), "lines": len(d.get("lines", []))}
                for d in payload.get("demands", payload.get("results", []))
            ],
        }
        history.insert(0, entry)
        history = history[:50]
        with open(os.path.join(_result_dir(), "history.json"), "w", encoding="utf-8") as f:
            json.dump(history, f, ensure_ascii=False, indent=2)
