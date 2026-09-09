#!/usr/bin/env python3
"""
财务凭证生成器 - 凭证生成模块
根据核算规则和业务数据生成会计凭证Excel文件
"""

import json
import re
from pathlib import Path
from typing import Dict, List, Any, Optional
from datetime import datetime
from dataclasses import dataclass, asdict, field

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


@dataclass
class VoucherEntry:
    """凭证分录"""
    entry_no: int                  # 分录序号
    summary: str                   # 摘要
    account_code: str              # 科目编码
    account_name: str              # 科目名称
    debit_amount: float = 0.0      # 借方金额
    credit_amount: float = 0.0     # 贷方金额
    auxiliary_type: str = ""       # 辅助核算类型
    auxiliary_code: str = ""       # 辅助核算编码
    auxiliary_name: str = ""       # 辅助核算名称
    currency: str = "人民币"        # 币种
    exchange_rate: float = 1.0     # 汇率
    original_amount: float = 0.0   # 原币金额


@dataclass
class Voucher:
    """会计凭证"""
    voucher_word: str = "记"       # 凭证字
    voucher_no: int = 1            # 凭证号
    voucher_date: str = ""         # 凭证日期
    attachment_count: int = 0      # 附单据数
    entries: List[VoucherEntry] = field(default_factory=list)
    total_debit: float = 0.0       # 借方合计
    total_credit: float = 0.0      # 贷方合计
    preparer: str = ""             # 制单人
    approver: str = ""             # 审核人
    poster: str = ""               # 过账人
    memo: str = ""                 # 备注
    ledger_code: str = ""          # 账簿编码


class VoucherGenerator:
    """凭证生成器"""

    # 支持的财务软件格式
    SUPPORTED_FORMATS = {
        'kingdee': '金蝶K/3 Cloud',
        'yonyou': '用友U8/U9',
        'sap': 'SAP',
        'oracle': 'Oracle Fusion'
    }

    def __init__(self, data_manager=None):
        self.data_manager = data_manager
        self.rules = []
        self.accounts = []

    def load_data(self, data_manager) -> None:
        """加载数据"""
        self.data_manager = data_manager
        self.rules = data_manager.get_rules_list()
        self.accounts = data_manager.get_account_codes()

    def _find_account(self, code_or_name: str) -> Optional[Dict]:
        """查找科目"""
        if not self.accounts:
            return None

        # 先按编码查找
        for account in self.accounts:
            if account.get('code') == code_or_name:
                return account

        # 再按名称查找
        for account in self.accounts:
            if account.get('name') == code_or_name:
                return account

        return None

    def _find_rule(self, business_type: str, source_field: str) -> Optional[Dict]:
        """查找核算规则"""
        for rule in self.rules:
            if (rule.get('business_type') == business_type or business_type == 'general') and \
               rule.get('source_field') == source_field:
                return rule
        return None

    def _match_auxiliary(self, value: str, auxiliary_type: str) -> Dict:
        """匹配辅助核算项"""
        if not self.data_manager:
            return {'code': '', 'name': value}

        value_lower = value.lower()

        if auxiliary_type == 'customer':
            # 匹配客户
            customers = self.data_manager.get_customers()
            for c in customers:
                if (c.get('name', '').lower() == value_lower or
                    c.get('short_name', '').lower() == value_lower or
                    c.get('code') == value):
                    return {'code': c.get('code', ''), 'name': c.get('name', '')}

        elif auxiliary_type == 'supplier':
            # 匹配供应商
            suppliers = self.data_manager.get_suppliers()
            for s in suppliers:
                if (s.get('name', '').lower() == value_lower or
                    s.get('short_name', '').lower() == value_lower or
                    s.get('code') == value):
                    return {'code': s.get('code', ''), 'name': s.get('name', '')}

        elif auxiliary_type == 'employee':
            # 匹配员工
            employees = self.data_manager.get_employees()
            for e in employees:
                if (e.get('name', '').lower() == value_lower or
                    e.get('code') == value):
                    return {'code': e.get('code', ''), 'name': e.get('name', '')}

        # 其他辅助资料
        aux_data = self.data_manager.get_auxiliary_items(auxiliary_type)
        for item in aux_data:
            if (item.get('name', '').lower() == value_lower or
                item.get('code') == value):
                return {'code': item.get('code', ''), 'name': item.get('name', '')}

        return {'code': '', 'name': value}

    def _get_exchange_rate(self, currency: str, date: str) -> float:
        """获取汇率"""
        if currency in ['人民币', 'CNY', 'RMB']:
            return 1.0

        if self.data_manager:
            rate = self.data_manager.get_exchange_rate(currency, date)
            if rate:
                return rate.get('rate_to_cny', 1.0)

        return 1.0

    def generate_vouchers(self,
                         source_data: pd.DataFrame,
                         business_type: str,
                         voucher_date: str,
                         ledger_code: str = "",
                         format_type: str = 'kingdee') -> Dict:
        """生成凭证

        Args:
            source_data: 源业务数据DataFrame
            business_type: 业务类型
            voucher_date: 凭证日期
            ledger_code: 账簿编码
            format_type: 目标财务软件格式

        Returns:
            生成结果
        """
        if source_data.empty:
            return {'success': False, 'error': '源数据为空'}

        vouchers = []
        unmapped_fields = []

        # 按行处理数据
        for idx, row in source_data.iterrows():
            entries = []

            # 遍历所有规则，生成借方分录
            debit_entry = VoucherEntry(
                entry_no=1,
                summary="",
                account_code="",
                account_name=""
            )

            credit_entry = VoucherEntry(
                entry_no=2,
                summary="",
                account_code="",
                account_name=""
            )

            row_mapped = False

            # 处理每一列
            for col_name in source_data.columns:
                value = row.get(col_name, '')

                # 查找匹配规则
                rule = self._find_rule(business_type, col_name)

                if rule:
                    row_mapped = True
                    target_account = self._find_account(rule.get('target_account_code', ''))

                    if target_account:
                        entry_type = rule.get('entry_type', 'debit')
                        amount = self._parse_amount(value)

                        if entry_type == 'debit':
                            debit_entry.entry_no = 1
                            debit_entry.summary = self._generate_summary(row, col_name, value)
                            debit_entry.account_code = target_account.get('code', '')
                            debit_entry.account_name = target_account.get('name', '')
                            debit_entry.debit_amount = amount

                            # 处理辅助核算
                            if target_account.get('is_auxiliary'):
                                for aux_type in target_account.get('auxiliary_types', []):
                                    aux_value = row.get(aux_type, value)
                                    aux_match = self._match_auxiliary(str(aux_value), aux_type)
                                    debit_entry.auxiliary_type = aux_type
                                    debit_entry.auxiliary_code = aux_match['code']
                                    debit_entry.auxiliary_name = aux_match['name']

                        elif entry_type == 'credit':
                            credit_entry.entry_no = 2
                            credit_entry.summary = self._generate_summary(row, col_name, value)
                            credit_entry.account_code = target_account.get('code', '')
                            credit_entry.account_name = target_account.get('name', '')
                            credit_entry.credit_amount = amount

                            # 处理辅助核算
                            if target_account.get('is_auxiliary'):
                                for aux_type in target_account.get('auxiliary_types', []):
                                    aux_value = row.get(aux_type, value)
                                    aux_match = self._match_auxiliary(str(aux_value), aux_type)
                                    credit_entry.auxiliary_type = aux_type
                                    credit_entry.auxiliary_code = aux_match['code']
                                    credit_entry.auxiliary_name = aux_match['name']

                else:
                    # 无法匹配的字段
                    if col_name not in unmapped_fields:
                        unmapped_fields.append(col_name)

            # 如果有借方或贷方分录，创建凭证
            if debit_entry.account_code or credit_entry.account_code:
                # 确保借贷平衡
                if debit_entry.debit_amount > 0 and credit_entry.credit_amount == 0:
                    credit_entry.credit_amount = debit_entry.debit_amount
                    # 尝试找默认的对方科目
                    default_credit = self._find_default_counter_account(business_type)
                    if default_credit:
                        credit_entry.account_code = default_credit.get('code', '')
                        credit_entry.account_name = default_credit.get('name', '')

                if credit_entry.credit_amount > 0 and debit_entry.debit_amount == 0:
                    debit_entry.debit_amount = credit_entry.credit_amount
                    # 尝试找默认的对方科目
                    default_debit = self._find_default_counter_account(business_type)
                    if default_debit:
                        debit_entry.account_code = default_debit.get('code', '')
                        debit_entry.account_name = default_debit.get('name', '')

                voucher = Voucher(
                    voucher_word="记",
                    voucher_no=idx + 1,
                    voucher_date=voucher_date,
                    attachment_count=1,
                    entries=[debit_entry, credit_entry],
                    total_debit=debit_entry.debit_amount,
                    total_credit=credit_entry.credit_amount,
                    ledger_code=ledger_code
                )
                vouchers.append(voucher)

        return {
            'success': True,
            'vouchers': vouchers,
            'unmapped_fields': unmapped_fields,
            'total_vouchers': len(vouchers)
        }

    def _parse_amount(self, value: Any) -> float:
        """解析金额"""
        if pd.isna(value):
            return 0.0

        if isinstance(value, (int, float)):
            return float(value)

        # 字符串金额处理
        value_str = str(value).strip()
        value_str = value_str.replace('¥', '').replace('元', '').replace(',', '').replace(' ', '')

        try:
            return float(value_str)
        except ValueError:
            return 0.0

    def _generate_summary(self, row: pd.Series, col_name: str, value: Any) -> str:
        """生成摘要"""
        # 从行数据中提取相关信息生成摘要
        summary_parts = []

        # 尝试从常用字段生成摘要
        for field_name in ['业务描述', '摘要', '说明', '交易类型', '类型']:
            if field_name in row.index and pd.notna(row[field_name]):
                summary_parts.append(str(row[field_name]))
                break

        if not summary_parts:
            summary_parts.append(f"{col_name}: {value}")

        return '/'.join(summary_parts[:2])

    def _find_default_counter_account(self, business_type: str) -> Optional[Dict]:
        """查找默认对方科目"""
        # 常见的对方科目
        default_accounts = {
            'bank_in': '1002',  # 银行存款
            'bank_out': '1002',
            'cash_in': '1001',  # 库存现金
            'cash_out': '1001',
            'receivable': '1122',  # 应收账款
            'payable': '2202',  # 应付账款
        }
        code = default_accounts.get(business_type)
        if code:
            return self._find_account(code)
        return None

    def to_excel(self,
                 vouchers: List[Voucher],
                 format_type: str = 'kingdee',
                 output_path: str = None) -> str:
        """导出为Excel文件

        Args:
            vouchers: 凭证列表
            format_type: 财务软件格式
            output_path: 输出路径

        Returns:
            输出文件路径
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "记账凭证"

        # 根据格式类型设置列标题
        headers = self._get_headers(format_type)

        # 写入表头
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')

        # 设置列宽
        self._set_column_widths(ws, format_type)

        # 写入数据
        current_row = 2
        for voucher in vouchers:
            for entry in voucher.entries:
                row_data = self._format_entry_row(entry, format_type)
                for col, value in enumerate(row_data, 1):
                    ws.cell(row=current_row, column=col, value=value)
                current_row += 1

            # 添加合计行
            ws.cell(row=current_row, column=1, value="")
            ws.cell(row=current_row, column=2, value="合计")
            ws.cell(row=current_row, column=5, value=voucher.total_debit)
            ws.cell(row=current_row, column=6, value=voucher.total_credit)
            current_row += 1

            # 添加空行分隔
            current_row += 1

        # 保存文件
        if output_path is None:
            SKILL_DIR = Path(__file__).parent.parent
            output_path = SKILL_DIR / "output"
            output_path.mkdir(exist_ok=True)
            output_path = output_path / f"vouchers_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

        wb.save(output_path)
        return str(output_path)

    def _get_headers(self, format_type: str) -> List[str]:
        """获取列标题"""
        base_headers = ['序号', '摘要', '科目编码', '科目名称', '借方金额', '贷方金额']

        if format_type == 'kingdee':
            return base_headers + ['辅助核算']
        elif format_type == 'yonyou':
            return base_headers + ['辅助核算', '项目', '部门']
        elif format_type == 'sap':
            return ['行号', '记账日期', '过账日期', '凭证类型', '科目', '科目描述',
                    '借方', '贷方', '货币', '汇率', '参考', '分配']
        elif format_type == 'oracle':
            return ['账户组合', '账户说明', '借方', '贷方', '币种', '辅助账户', '说明']

        return base_headers

    def _set_column_widths(self, ws, format_type: str) -> None:
        """设置列宽"""
        widths = {
            'kingdee': [6, 30, 12, 20, 15, 15, 20],
            'yonyou': [6, 30, 12, 20, 15, 15, 15, 10, 10],
            'sap': [6, 12, 12, 6, 10, 30, 15, 15, 5, 10, 15, 15],
            'oracle': [15, 30, 15, 15, 8, 15, 30]
        }

        width_list = widths.get(format_type, widths['kingdee'])
        for i, width in enumerate(width_list, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

    def _format_entry_row(self, entry: VoucherEntry, format_type: str) -> List:
        """格式化分录行"""
        base_row = [
            entry.entry_no,
            entry.summary,
            entry.account_code,
            entry.account_name,
            entry.debit_amount if entry.debit_amount else "",
            entry.credit_amount if entry.credit_amount else ""
        ]

        if format_type == 'kingdee':
            aux_text = f"{entry.auxiliary_type}:{entry.auxiliary_name}" if entry.auxiliary_type else ""
            return base_row + [aux_text]
        elif format_type == 'yonyou':
            return base_row + [entry.auxiliary_name, "", ""]
        elif format_type == 'sap':
            return [entry.entry_no, "", "", "SA", entry.account_code, entry.account_name,
                   entry.debit_amount, entry.credit_amount, entry.currency,
                   entry.exchange_rate, "", ""]
        elif format_type == 'oracle':
            return [entry.account_code, entry.account_name,
                   entry.debit_amount, entry.credit_amount, entry.currency,
                   entry.auxiliary_code, entry.summary]

        return base_row

    def generate_template(self, format_type: str = 'kingdee', output_path: str = None) -> str:
        """生成空白凭证模板"""
        wb = Workbook()
        ws = wb.active
        ws.title = "凭证导入模板"

        headers = self._get_headers(format_type)
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')

        # 添加示例行
        example_data = [
            [1, "示例摘要", "1001", "库存现金", 1000.00, ""],
            [2, "", "1002", "银行存款", "", 1000.00]
        ]

        for row_idx, row_data in enumerate(example_data, 2):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        self._set_column_widths(ws, format_type)

        if output_path is None:
            SKILL_DIR = Path(__file__).parent.parent
            templates_dir = SKILL_DIR / "assets"
            templates_dir.mkdir(exist_ok=True)
            output_path = templates_dir / f"{format_type}_voucher_template.xlsx"

        wb.save(output_path)
        return str(output_path)


def get_voucher_generator(data_manager=None) -> VoucherGenerator:
    """获取凭证生成器实例"""
    return VoucherGenerator(data_manager)


if __name__ == '__main__':
    # 测试代码
    generator = VoucherGenerator()
    template_path = generator.generate_template('kingdee')
    print(f"模板已生成: {template_path}")
